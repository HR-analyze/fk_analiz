"""Выгрузки раздела «КЦ»: Excel по виджетам и PDF-сводка.

Берут уже посчитанный kc_aggregate.build_summary — тот же, что видит страница, поэтому
цифры в файлах и на экране совпадают при одинаковых фильтрах. В каждом файле — оговорка,
что сумма штук по операциям не является уникальным готовым выпуском.
"""
import io
from datetime import datetime

from exports import TZ, _bar_drawing, _pie_drawing, _register_fonts
from kc_aggregate import MIX_WARNING, NORM_NOTE, PCS_NOTE, change_text, kpi_rows

NO_DATA = "нет данных"


def _v(value):
    return "" if value is None else value


def _pct(value, total):
    return round(value / total * 100, 1) if total else ""


def _group_rows(rows, key):
    return [[r[key], r["completed"], r["pcs"], r["person_hours"], _v(r["avg_duration"]), _v(r["avg_people"]),
             _v(r["pcs_per_hour"]), _v(r["pcs_per_ph"])] for r in rows]


GROUP_COLS = ["Завершено операций", "Обработано изделий, шт", "Человеко-часы", "Ср. длительность, мин",
              "Ср. людей", "шт/ч", "шт/чел·ч"]


def widget_rows(summary, widget):
    """(заголовок листа, колонки, строки)."""
    k = summary["kpi"]
    lfl = summary.get("lfl") or {}
    if widget == "kpi":
        prev_col = f"Прошлый период {lfl['prev_label']}" if lfl.get("available") else "Прошлый период"
        rows = []
        for title, text, key in kpi_rows(k):
            c = (lfl.get("kpi") or {}).get(key) if key and lfl.get("available") else None
            prev = "" if not c else (NO_DATA if c["prev"] is None else c["prev"])
            rows.append([title, text, prev, "" if not c else _v(c["change"])])
        rows += [["Начато операций", k["started"], "", ""], ["Отменено (из начатых)", k["cancelled"], "", ""],
                 ["Средняя выработка, шт/ч", _v(k["pcs_per_hour"]), "", ""],
                 ["Среднее число людей в операции", _v(k["avg_people"]), "", ""],
                 ["Операций с некорректным временем", k["invalid"], "", ""],
                 ["Внимание: штуки", PCS_NOTE, "", ""], ["Внимание: выработка по цеху", MIX_WARNING, "", ""]]
        return "KPI", ["Показатель", "Значение", prev_col, "Изменение, %"], rows
    if widget == "daily":
        return "По дням", ["Дата", "День", "Начато", "Завершено", "Обработано изделий, шт", "Человеко-часы",
                           "Ср. длительность, мин", "шт/чел·ч", "Остановок", "Простои, мин",
                           "Неделей раньше, дата", "Неделей раньше, шт", "К неделе раньше, %"], [
            [d["date"], d["weekday"], d["started"], d["completed"], d["pcs"], d["person_hours"], _v(d["avg_duration"]),
             _v(d["pcs_per_ph"]), d["pauses"], d["stop"], d.get("lfl_date", ""), _v(d.get("lfl_pcs")),
             _v(d.get("lfl_change"))] for d in summary["daily"]]
    if widget == "hourly":
        return "По часам", ["Час завершения", "Обработано изделий, шт", "Завершено операций"], [
            [h["label"], h["value"], h["completed"]] for h in summary["hourly"]]
    if widget == "products":
        return "По продукции", ["Продукция"] + GROUP_COLS, _group_rows(summary["products"], "product")
    if widget == "operations":
        return "По операциям", ["Операция"] + GROUP_COLS, _group_rows(summary["operations"], "operation")
    if widget == "norms":
        prev = f"шт/чел·ч, {lfl['prev_label']}" if lfl.get("available") else "шт/чел·ч, прошлый период"
        return "Нормы", ["Продукция", "Операция", "Выполнений", "Обработано изделий, шт", "Ср. время, мин",
                         "Мин. время, мин", "Макс. время, мин", "Ср. людей", "Человеко-часы", "шт/ч", "шт/чел·ч",
                         prev, "Изменение, %", "Некорректных записей"], [
            [r["product"], r["operation"], r["completed"], r["pcs"], _v(r["avg_duration"]), _v(r["min_duration"]),
             _v(r["max_duration"]), _v(r["avg_people"]), r["person_hours"], _v(r["pcs_per_hour"]), _v(r["pcs_per_ph"]),
             _v(r["prev_pcs_per_ph"]), _v(r["pcs_per_ph_change"]), r["invalid"]] for r in summary["norms"]]
    if widget == "shifts":
        return "По сменам", ["Дата смены", "Смена", "Завершено операций", "Обработано изделий, шт", "Человеко-часы",
                             "шт/чел·ч", "Через полночь"], [
            [r["shift_date"], r["shift_name"], r["completed"], r["pcs"], r["person_hours"], _v(r["pcs_per_ph"]),
             r["midnight"]] for r in summary["shifts"]]
    if widget == "employees":
        return "Сотрудники", ["Сотрудник", "Начал операций", "Завершил", "Отменил", "Операций с его отметкой",
                              "Шт в завершённых им операциях (группы, не личная выработка)", "Остановок",
                              "Простои, мин"], [
            [e["employee"], e["started"], e["closed"], e["cancelled"], e["ops"], e["pcs_closed"], e["pauses"], e["stop"]]
            for e in summary["employees"]]
    if widget == "activity":
        return "Активность по часам", ["Сотрудник"] + [f"{h:02d}" for h in range(24)], [
            [e["employee"]] + e["hours"] for e in summary["employees"]]
    if widget == "reasons":
        total = sum(r["value"] for r in summary["pause_reasons"])
        return "Остановки по причинам", ["Причина", "Раз", "Минуты", "Доля, %"], [
            [r["label"], r["count"], r["value"], _pct(r["value"], total)] for r in summary["pause_reasons"]]
    if widget == "pauses":
        return "Критические остановки", ["Дата", "Причина", "Начало", "Конец", "Минуты", "Сотрудник"], [
            [p["date"], p["reason"], p["start_time"], p["end_time"], p["minutes"], p["user_name"]]
            for p in summary["pause_rows"]]
    if widget == "open_ops":
        return "В работе", ["Продукция", "Операция", "Начало", "Открыта, мин", "Людей", "Начал"], [
            [r["product"], r["operation"], f"{r['start_date']} {r['start_time']}", _v(r["open_min"]), _v(r["people"]),
             r["started_by"]] for r in summary["open_ops"]]
    if widget == "ops":
        return "Операции", ["Дата", "Продукция", "Операция", "Начало", "Завершение", "Длительность, мин", "Шт",
                            "Людей", "Человеко-часы", "шт/ч", "шт/чел·ч", "Начал", "Завершил", "Статус"], [
            [r["date"], r["product"], r["operation"], f"{r['start_date']} {r['start_time']}",
             f"{r['end_date']} {r['end_time']}" if r["end_time"] else "",
             _v(r["duration_min"]) if r["valid"] or r["duration_min"] is None else f"{r['duration_min']} (некорр.)",
             _v(r["pcs"]), _v(r["people"]), _v(r["person_hours"]), _v(r["pcs_per_hour"]), _v(r["pcs_per_ph"]),
             r["started_by"], r["closed_by"], r["status_name"]] for r in summary["ops"]]
    raise KeyError(widget)


ALL_WIDGETS = ("kpi", "daily", "hourly", "products", "operations", "norms", "shifts", "employees", "activity",
               "reasons", "pauses", "open_ops", "ops")


def build_xlsx(summary, widget, filters):
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    widgets = ALL_WIDGETS if widget == "all" else [widget]
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    head_fill = PatternFill("solid", fgColor="E8EEFB")
    head_font = Font(bold=True)
    for name in widgets:
        title, columns, rows = widget_rows(summary, name)
        ws = wb.create_sheet(title[:31])
        ws.append(columns)
        for row in rows:
            ws.append(row)
        if not rows:
            ws.append([NO_DATA])
        for idx, column in enumerate(columns, 1):
            cell = ws.cell(row=1, column=idx)
            cell.fill, cell.font = head_fill, head_font
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            longest = max([len(str(column))] + [len(str(r[idx - 1])) for r in rows[:200] if idx - 1 < len(r)] or [0])
            ws.column_dimensions[get_column_letter(idx)].width = min(max(longest + 2, 9), 46)
        ws.freeze_panes = "A2"
        if rows:
            ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{len(rows) + 1}"
    ws = wb.create_sheet("Фильтры")
    ws.append(["Параметр", "Значение"])
    for name, value in filters:
        ws.append([name, value])
    ws.append(["Выгружено", datetime.now(TZ).strftime("%d.%m.%Y %H:%M")])
    ws.append([])
    ws.append(["Обработано изделий", PCS_NOTE])
    ws.append(["Выработка по цеху", MIX_WARNING])
    ws.append(["Нормы", NORM_NOTE])
    ws.append(["Сотрудники", "people — численность группы; выработка конкретным людям не приписывается"])
    if summary.get("ops_total", 0) > len(summary.get("ops", [])):
        ws.append(["Операции", f"в листе первые {len(summary['ops'])} из {summary['ops_total']} — сузь период"])
    for idx in (1, 2):
        ws.cell(row=1, column=idx).fill = head_fill
        ws.cell(row=1, column=idx).font = head_font
        ws.column_dimensions[get_column_letter(idx)].width = 40 if idx == 1 else 90
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def build_pdf(summary, analysis, filters, title="Сводка: кондитерский цех"):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import KeepTogether, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    font, bold = _register_fonts()
    out = io.BytesIO()
    doc = SimpleDocTemplate(out, pagesize=landscape(A4), leftMargin=12 * mm, rightMargin=12 * mm,
                            topMargin=12 * mm, bottomMargin=12 * mm, title=title)
    width = doc.width
    h1 = ParagraphStyle("h1", fontName=bold, fontSize=17, leading=21, spaceAfter=4)
    h2 = ParagraphStyle("h2", fontName=bold, fontSize=11, leading=14, spaceBefore=8, spaceAfter=4)
    body = ParagraphStyle("body", fontName=font, fontSize=8.5, leading=12)
    muted = ParagraphStyle("muted", fontName=font, fontSize=8, leading=11, textColor=colors.HexColor("#667085"))
    warn = ParagraphStyle("warn", fontName=font, fontSize=8, leading=11, textColor=colors.HexColor("#8A5A08"),
                          backColor=colors.HexColor("#FFF7E8"), borderPadding=4)
    cell = ParagraphStyle("c", fontName=font, fontSize=6.6, leading=8)

    def table(columns, rows, limit=40, size=6.8):
        data = [[Paragraph(str(c), ParagraphStyle("h", fontName=bold, fontSize=size, leading=size + 1.4)) for c in columns]]
        data += [[Paragraph(str(c), cell) for c in row] for row in rows[:limit]]
        t = Table(data, repeatRows=1)
        t.setStyle(TableStyle([
            ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#E6EAF0")),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#EFF3FB")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("TOPPADDING", (0, 0), (-1, -1), 2.2), ("BOTTOMPADDING", (0, 0), (-1, -1), 2.2),
            ("LEFTPADDING", (0, 0), (-1, -1), 3), ("RIGHTPADDING", (0, 0), (-1, -1), 3),
        ]))
        return t

    story = [Paragraph(title, h1),
             Paragraph(", ".join(f"{n}: {v}" for n, v in filters if v) or "фильтры не заданы", muted),
             Paragraph(f"Сформировано {datetime.now(TZ).strftime('%d.%m.%Y %H:%M')}", muted), Spacer(1, 4 * mm)]
    lfl = summary.get("lfl") or {}

    def tile(label, value, key):
        delta = ""
        c = (lfl.get("kpi") or {}).get(key) if key else None
        if lfl.get("available") and not lfl.get("prev_empty") and c and c["change"] is not None:
            delta = f"<br/><font size=6.5 color='#667085'>{change_text(c['change'])} к {lfl['prev_label']}</font>"
        return Paragraph(f"<font size=7 color='#667085'>{label}</font><br/><font size=12 name='{bold}'>{value}</font>{delta}",
                         ParagraphStyle("k", fontName=font, leading=16))

    tiles = [tile(*row) for row in kpi_rows(summary["kpi"])]
    grid = Table([tiles[i:i + 4] for i in (0, 4)], colWidths=[width / 4] * 4)
    grid.setStyle(TableStyle([("BOX", (0, 0), (-1, -1), 0.4, colors.HexColor("#E6EAF0")),
                              ("INNERGRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#E6EAF0")),
                              ("VALIGN", (0, 0), (-1, -1), "TOP"),
                              ("TOPPADDING", (0, 0), (-1, -1), 6), ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                              ("LEFTPADDING", (0, 0), (-1, -1), 7)]))
    story += [grid, Spacer(1, 3 * mm), Paragraph("⚠ " + PCS_NOTE.capitalize() + ". " + MIX_WARNING, warn), Spacer(1, 4 * mm)]
    for heading, lines in analysis:
        story.append(KeepTogether([Paragraph(heading, h2)] + [Paragraph(line, body) for line in lines]))
    story.append(PageBreak())

    charts = [
        ("Завершено операций по дням", [{"label": d["date"][5:], "value": d["completed"]} for d in summary["daily"]], "#3B6EF5"),
        ("Обработано изделий по дням, шт (сумма по операциям)", [{"label": d["date"][5:], "value": d["pcs"]} for d in summary["daily"]], "#7C5CF5"),
        ("Обработано изделий по часу завершения, шт", [{"label": h["label"], "value": h["value"]} for h in summary["hourly"]], "#3FBF94"),
        ("Человеко-часы по продукции", summary["labor_products"][:10], "#F0A63C"),
        ("Человеко-часы по операциям", summary["labor_operations"][:10], "#2A78D6"),
    ]
    story.append(Paragraph("Графики", h2))
    for name, items, color_hex in charts:
        story.append(KeepTogether([Paragraph(name, body), _bar_drawing(items, width, 120, font, color_hex), Spacer(1, 3 * mm)]))
    palette = ["#2A78D6", "#EB6834", "#1BAF7A", "#EDA100", "#E87BA4", "#008300", "#4A3AA7", "#E34948"]
    story.append(KeepTogether([Paragraph("Критические остановки по причинам, мин", body),
                               _pie_drawing(summary["pause_reasons"], width, 116, font, palette)]))
    story.append(PageBreak())

    story.append(Paragraph("Нормы по продукции и операциям", h2))
    story.append(Paragraph(NORM_NOTE, muted))
    for widget, limit in (("norms", 60), ("products", 30), ("operations", 30), ("daily", 40), ("shifts", 30),
                          ("employees", 30), ("reasons", 15), ("pauses", 40), ("open_ops", 30), ("ops", 120)):
        heading, columns, data = widget_rows(summary, widget)
        if not data:
            continue
        if widget != "norms":
            story.append(Paragraph(heading, h2))
        story.append(table(columns, data, limit=limit))
        if len(data) > limit:
            story.append(Paragraph(f"Показаны первые {limit} из {len(data)} строк; полный список — в Excel.", muted))
        story.append(Spacer(1, 4 * mm))
    doc.build(story)
    return out.getvalue()
