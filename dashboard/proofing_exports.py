"""Выгрузки раздела «Расстойка»: Excel по виджетам и PDF-сводка.

Берут уже посчитанный proofing_aggregate.build_summary — тот же, что видит страница,
поэтому цифры в файлах и на экране совпадают при одинаковых фильтрах.
exports.py (ФК/БХМ) не меняется: оттуда берутся только шрифты и рисование графиков.
"""
import io
from datetime import datetime

from exports import TZ, _bar_drawing, _pie_drawing, _register_fonts
from proofing_aggregate import UNIT_SHORT, change_text, kpi_rows

NO_DATA = "нет данных"


def _v(value):
    """Пустая ячейка вместо None: в Excel по ней можно фильтровать и считать."""
    return "" if value is None else value


def _pct(value, total):
    return round(value / total * 100, 1) if total else ""


def _stage(s, unit_default="racks"):
    if not s:
        return ["", "", "", "", "", "", ""]
    unit = UNIT_SHORT.get(s["rack_unit"] or unit_default, "шп.")
    load = "" if s["racks_start"] is None else f"{s['racks_start']}" + (
        f"→{s['racks_end']}" if s["racks_end"] not in (None, s["racks_start"]) else "") + f" {unit}"
    period = f"{s['start_date']} {s['start_time']}"
    end = f"{s['end_date']} {s['end_time']}" if s["end"] else ("отменён" if s["status"] == "cancelled" else "не закрыт")
    return [period, end, _v(s["duration_min"]) if s["valid"] else ("" if s["duration_min"] is None else f"{s['duration_min']} (некорр.)"),
            load, s["started_by"], s["closed_by"], s["equipment"]]


def widget_rows(summary, widget):
    """(заголовок листа, колонки, строки)."""
    k = summary["kpi"]
    lfl = summary.get("lfl") or {}
    if widget == "kpi":
        prev_col = f"Прошлый период {lfl['prev_label']}" if lfl.get("available") else "Прошлый период"
        rows = []
        for title, text, key in kpi_rows(k):
            c = (lfl.get("kpi") or {}).get(key) if key and lfl.get("available") else None
            prev = "" if not c else (NO_DATA if c["prev"] is None else c["prev"])
            rows.append([title, text, prev, "" if not c else _v(c["change"])])
        rows += [["Отменено партий (из начатых)", k["cancelled"], "", ""],
                 ["Этапов расстойки в среднем", k["proofing_count"], "", ""],
                 ["Этапов отпекания в среднем", k["baking_count"], "", ""],
                 ["Этапов с некорректным временем", k["invalid_stages"], "", ""],
                 ["Остановки, мин", k["stop_minutes"], "", ""]]
        return "KPI", ["Показатель", "Значение", prev_col, "Изменение, %"], rows
    if widget == "daily":
        return "Выпуск по дням", ["Дата", "День", "Начато партий", "Завершено партий", "Выпуск, шт",
                                  "Ср. расстойка, мин", "Ср. отпекание, мин", "Отпекание, шт/ч", "Остановок",
                                  "Остановки, мин", "Неделей раньше, дата", "Неделей раньше, шт", "К неделе раньше, %"], [
            [d["date"], d["weekday"], d["started"], d["completed"], d["output"], _v(d["avg_proofing"]),
             _v(d["avg_baking"]), _v(d["rate"]), d["pauses"], d["stop"], d.get("lfl_date", ""),
             _v(d.get("lfl_output")), _v(d.get("lfl_change"))] for d in summary["daily"]]
    if widget == "hourly":
        return "Выпуск по часам", ["Час окончания отпекания", "Выпуск, шт", "Партий", "В среднем за день, шт"], [
            [h["label"], h["value"], h["batches"], h["per_day"]] for h in summary["hourly"]]
    if widget == "ovens":
        return "По печам", ["Печь", "Единица загрузки", "Партий", "Выпуск, шт", "Загружено", "Ср. отпекание, мин",
                            "Отпекание, шт/ч", "Занято, ч", "Занято, % периода", "Одновременно до, партий",
                            "Остановок", "Остановки, мин"], [
            [o["name"], o["unit_name"], o["batches"], o["output"], o["loaded_text"], _v(o["avg_baking"]), _v(o["rate"]),
             o["busy_hours"], _v(o["utilization"]), o["max_parallel"], o["pauses"], o["stop"]] for o in summary["ovens"]]
    if widget == "products":
        return "По изделиям", ["Изделие", "Начато", "Завершено", "Выпуск, шт", "Ср. расстойка, мин",
                               "Ср. отпекание, мин", "Ср. цикл, мин", "Отпекание, шт/ч"], [
            [p["product"], p["started"], p["completed"], p["output"], _v(p["avg_proofing"]), _v(p["avg_baking"]),
             _v(p["avg_cycle"]), _v(p["rate"])] for p in summary["products"]]
    if widget == "employees":
        return "Сотрудники", ["Сотрудник", "Партий", "Отметок", "Начал расстойку", "Закрыл расстойку",
                              "Начал отпекание", "Закрыл отпекание", "Выпуск закрытых отпеканий, шт",
                              "Остановок", "Остановки, мин"], [
            [e["employee"], e["batches"], e["marks"], e["start_proofing"], e["end_proofing"], e["start_baking"],
             e["end_baking"], e["output_closed"], e["pauses"], e["stop"]] for e in summary["employees"]]
    if widget == "load":
        rows = []
        for r in summary["load"]["rows"]:
            for h in r["hours"]:
                rows.append([r["name"], f"{h['hour']:02d}:00", h["busy_min"], h["pct"], h["parallel"]])
        return "Загрузка по часам", ["Оборудование", "Час", "Занято в среднем, мин/ч", "Занято, %",
                                     "Партий одновременно в среднем"], rows
    if widget == "load_summary":
        return "Загрузка оборудования", ["Оборудование", "Этапов", "Занято, ч (без пересечений)",
                                         "Сумма по партиям, ч", "Занято, % периода", "Одновременно до, партий"], [
            [r["name"], r["stages"], r["busy_hours"], r["sum_hours"], _v(r["utilization"]), r["max_parallel"]]
            for r in summary["load"]["rows"]]
    if widget == "reasons":
        total = sum(r["value"] for r in summary["pause_reasons"])
        return "Остановки по причинам", ["Причина", "Раз", "Минуты", "Доля, %"], [
            [r["label"], r["count"], r["value"], _pct(r["value"], total)] for r in summary["pause_reasons"]]
    if widget == "pauses":
        return "Критические остановки", ["Дата", "Оборудование", "Причина", "Начало", "Конец", "Минуты", "Сотрудник"], [
            [p["date"], p["equipment"], p["reason"], p["start_time"], p["end_time"], p["minutes"], p["user_name"]]
            for p in summary["pause_rows"]]
    if widget == "open_stages":
        return "Незакрытые этапы", ["Изделие", "Этап", "Оборудование", "Начало", "Открыт, мин", "Дольше лимита",
                                    "Начал", "Следующая отметка"], [
            [r["product"], r["stage"], r["equipment"], f"{r['start_date']} {r['start_time']}", r["open_min"],
             "да" if r["stale"] else "", r["started_by"], r["next_mark"]] for r in summary["open_stages"]]
    if widget == "batches":
        return "Партии", ["Дата", "Старт", "Изделие", "Статус", "Печь",
                          "Расстойка: начало", "Расстойка: конец", "Расстойка, мин", "Расстойка: загрузка",
                          "Расстойка: начал", "Расстойка: закончил", "Расстойка: оборудование",
                          "Отпекание: начало", "Отпекание: конец", "Отпекание, мин", "Отпекание: загрузка",
                          "Отпекание: начал", "Отпекание: закончил", "Отпекание: оборудование",
                          "Вход, шт", "Выход, шт", "Отпекание, шт/ч", "Цикл, мин", "Сотрудники", "Фото", "ID партии"], [
            [b["date"], b["start_time"], b["product"], b["status_name"], b["oven_name"]]
            + _stage(b["proofing"]) + _stage(b["baking"], b["oven_unit"] or "racks")
            + [_v(b["pieces_in"]), _v(b["pieces_out"]), _v(b["rate"]), _v(b["cycle_min"]), ", ".join(b["employees"]),
               b["photo_count"], b["id"]] for b in summary["batches"]]
    raise KeyError(widget)


ALL_WIDGETS = ("kpi", "daily", "hourly", "ovens", "products", "employees", "load_summary", "load", "reasons",
               "pauses", "open_stages", "batches")


def build_xlsx(summary, widget, filters):
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    widgets = ALL_WIDGETS if widget == "all" else [widget]
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    head_fill = PatternFill("solid", fgColor="E8EEFB")
    head_font = Font(bold=True)
    for name in widgets:
        title, columns, rows = widget_rows(summary, name)
        ws = wb.create_sheet(title[:31])
        ws.append(columns)
        for row in rows:
            ws.append(row)
        if not rows:
            ws.append([NO_DATA])
        for idx, column in enumerate(columns, 1):
            cell = ws.cell(row=1, column=idx)
            cell.fill, cell.font = head_fill, head_font
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            longest = max([len(str(column))] + [len(str(r[idx - 1])) for r in rows[:200] if idx - 1 < len(r)] or [0])
            ws.column_dimensions[get_column_letter(idx)].width = min(max(longest + 2, 11), 46)
        ws.freeze_panes = "A2"
        if rows:
            ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{len(rows) + 1}"
    ws = wb.create_sheet("Фильтры")
    ws.append(["Параметр", "Значение"])
    for name, value in filters:
        ws.append([name, value])
    ws.append(["Выгружено", datetime.now(TZ).strftime("%d.%m.%Y %H:%M")])
    ws.append([])
    ws.append(["Выпуск", "по концу отпекания; отменённые партии не учитываются"])
    ws.append(["Поды и шпильки", "считаются раздельно, по единице каждой печи"])
    if summary.get("batches_total", 0) > len(summary.get("batches", [])):
        ws.append(["Партии", f"в листе первые {len(summary['batches'])} из {summary['batches_total']} — сузь период"])
    for idx in (1, 2):
        ws.cell(row=1, column=idx).fill = head_fill
        ws.cell(row=1, column=idx).font = head_font
        ws.column_dimensions[get_column_letter(idx)].width = 40
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def build_pdf(summary, analysis, filters, title="Сводка: расстойка и отпекание"):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import KeepTogether, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    font, bold = _register_fonts()
    out = io.BytesIO()
    doc = SimpleDocTemplate(out, pagesize=A4, leftMargin=13 * mm, rightMargin=13 * mm,
                            topMargin=13 * mm, bottomMargin=13 * mm, title=title)
    width = doc.width
    h1 = ParagraphStyle("h1", fontName=bold, fontSize=17, leading=21, spaceAfter=4)
    h2 = ParagraphStyle("h2", fontName=bold, fontSize=11, leading=14, spaceBefore=8, spaceAfter=4)
    body = ParagraphStyle("body", fontName=font, fontSize=8.5, leading=12)
    muted = ParagraphStyle("muted", fontName=font, fontSize=8, leading=11, textColor=colors.HexColor("#667085"))
    cell = ParagraphStyle("c", fontName=font, fontSize=6.4, leading=7.8)

    def table(columns, rows, widths=None, limit=40, size=6.8):
        data = [[Paragraph(str(c), ParagraphStyle("h", fontName=bold, fontSize=size, leading=size + 1.4)) for c in columns]]
        data += [[Paragraph(str(c), cell) for c in row] for row in rows[:limit]]
        t = Table(data, colWidths=widths, repeatRows=1)
        t.setStyle(TableStyle([
            ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#E6EAF0")),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#EFF3FB")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("TOPPADDING", (0, 0), (-1, -1), 2.2), ("BOTTOMPADDING", (0, 0), (-1, -1), 2.2),
            ("LEFTPADDING", (0, 0), (-1, -1), 3), ("RIGHTPADDING", (0, 0), (-1, -1), 3),
        ]))
        return t

    story = [Paragraph(title, h1)]
    story.append(Paragraph(", ".join(f"{n}: {v}" for n, v in filters if v) or "фильтры не заданы", muted))
    story.append(Paragraph(f"Сформировано {datetime.now(TZ).strftime('%d.%m.%Y %H:%M')}", muted))
    story.append(Spacer(1, 5 * mm))

    lfl = summary.get("lfl") or {}

    def tile(label, value, key):
        delta = ""
        c = (lfl.get("kpi") or {}).get(key) if key else None
        if lfl.get("available") and not lfl.get("prev_empty") and c and c["change"] is not None:
            delta = f"<br/><font size=6.5 color='#667085'>{change_text(c['change'])} к {lfl['prev_label']}</font>"
        return Paragraph(f"<font size=7 color='#667085'>{label}</font><br/><font size=12 name='{bold}'>{value}</font>{delta}",
                         ParagraphStyle("k", fontName=font, leading=16))

    tiles = [tile(*row) for row in kpi_rows(summary["kpi"])]
    grid = Table([tiles[i:i + 4] for i in (0, 4)], colWidths=[width / 4] * 4)
    grid.setStyle(TableStyle([("BOX", (0, 0), (-1, -1), 0.4, colors.HexColor("#E6EAF0")),
                              ("INNERGRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#E6EAF0")),
                              ("VALIGN", (0, 0), (-1, -1), "TOP"),
                              ("TOPPADDING", (0, 0), (-1, -1), 6), ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                              ("LEFTPADDING", (0, 0), (-1, -1), 7)]))
    story += [grid, Spacer(1, 5 * mm)]
    for heading, lines in analysis:
        story.append(KeepTogether([Paragraph(heading, h2)] + [Paragraph(line, body) for line in lines]))
    story.append(PageBreak())

    charts = [
        ("Выпуск по дням (по концу отпекания), шт", [{"label": d["date"][5:], "value": d["output"]} for d in summary["daily"]], "#3B6EF5"),
        ("Выпуск по часам окончания отпекания, шт", [{"label": h["label"], "value": h["value"]} for h in summary["hourly"]], "#7C5CF5"),
        ("Производительность отпекания по печам, шт/ч", [{"label": o["name"], "value": o["rate"] or 0} for o in summary["ovens"]], "#3FBF94"),
        ("Топ изделий по выпуску, шт", summary["top_products"][:10], "#F0A63C"),
        ("Партии по сотрудникам", summary["top_employees"][:10], "#2A78D6"),
    ]
    story.append(Paragraph("Графики", h2))
    for name, items, color_hex in charts:
        story.append(KeepTogether([Paragraph(name, body), _bar_drawing(items, width, 128, font, color_hex), Spacer(1, 3 * mm)]))
    palette = ["#2A78D6", "#EB6834", "#1BAF7A", "#EDA100", "#E87BA4", "#008300", "#4A3AA7", "#E34948"]
    story.append(KeepTogether([Paragraph("Критические остановки по причинам, мин", body),
                               _pie_drawing(summary["pause_reasons"], width, 120, font, palette)]))
    story.append(PageBreak())

    # Загрузка: строка — оборудование, столбцы — часы, значение — % занятости часа.
    load = summary["load"]
    story.append(Paragraph(f"Загрузка по часам, % времени часа (в среднем за {load['days']} дн., пересечения партий не суммируются)", h2))
    head = ["Оборудование"] + [f"{h:02d}" for h in range(24)]
    rows = [[r["name"]] + [("" if h["pct"] == 0 else f"{h['pct']:.0f}") for h in r["hours"]] for r in load["rows"]]
    story.append(table(head, rows, widths=[30 * mm] + [(width - 30 * mm) / 24] * 24, size=6))
    story.append(Spacer(1, 4 * mm))

    for widget, limit in (("load_summary", 10), ("ovens", 10), ("daily", 40), ("products", 30),
                          ("employees", 20), ("reasons", 15), ("open_stages", 30)):
        heading, columns, data = widget_rows(summary, widget)
        if not data:
            continue
        story.append(Paragraph(heading, h2))
        story.append(table(columns, data, limit=limit))
        if len(data) > limit:
            story.append(Paragraph(f"Показаны первые {limit} из {len(data)} строк.", muted))
        story.append(Spacer(1, 4 * mm))

    batches = summary["batches"]
    if batches:
        story.append(Paragraph("Партии", h2))
        rows = []
        for b in batches[:80]:
            p, k = b["proofing"] or {}, b["baking"] or {}
            rows.append([f"{b['date'][5:]} {b['start_time']}", b["product"], b["status_name"], b["oven_name"] or "—",
                         f"{p.get('start_time', '')}–{p.get('end_time') or '…'}" if p else "—",
                         _v(p.get("duration_min")), f"{k.get('start_time', '')}–{k.get('end_time') or '…'}" if k else "—",
                         _v(k.get("duration_min")), _v(b["pieces_in"]), _v(b["pieces_out"]), _v(b["rate"]),
                         ", ".join(b["employees"])])
        story.append(table(["Старт", "Изделие", "Статус", "Печь", "Расстойка", "мин", "Отпекание", "мин",
                            "Вход, шт", "Выход, шт", "шт/ч", "Сотрудники"], rows, limit=80,
                           widths=[16 * mm, 34 * mm, 15 * mm, 17 * mm, 17 * mm, 8 * mm, 17 * mm, 8 * mm,
                                   12 * mm, 12 * mm, 10 * mm, width - 166 * mm]))
        if summary.get("batches_total", 0) > 80:
            story.append(Paragraph(f"Показаны первые 80 из {summary['batches_total']} партий; полный список — в Excel.", muted))
    doc.build(story)
    return out.getvalue()
