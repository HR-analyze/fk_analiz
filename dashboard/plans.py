"""Планы производства: разбор Excel-файла и хранение на диске.

К базе бота отношения не имеет — планы лежат отдельным JSON-файлом на диске ВМ.
"""
import json
import logging
import os
import re
import tempfile
from datetime import date, datetime
from pathlib import Path
from threading import Lock

log = logging.getLogger("fk.plans")

PLAN_FILE = Path(os.getenv("PLAN_FILE", "/app/data/plans.json"))
PLAN_UPLOAD_TOKEN = os.getenv("PLAN_UPLOAD_TOKEN", "").strip()
MAX_UPLOAD_BYTES = int(os.getenv("PLAN_MAX_UPLOAD_BYTES", str(5 * 1024 * 1024)))
MAX_ROWS = int(os.getenv("PLAN_MAX_ROWS", "20000"))

# Заголовки колонок рабочего шаблона. Веса, а не первое попадание: в файле есть и
# «Код номенклатуры», и «Сокращенное название» — по одной подстроке их не различить.
COLUMN_HINTS = {
    "product": [("сокращен", 10), ("наименование", 8), ("продукт", 8),
                ("номенклатур", 3), ("код", -30), ("точная", -10)],
    "equipment": [("оборудован", 10), ("линия", 5)],
    "quantity": [("план количество", 12), ("количество", 8), ("план, шт", 8),
                 ("план шт", 8), ("факт", -30), ("дата", -30), ("план", 4)],
    "date": [("дата план", 12), ("дата", 8), ("код", -30)],
}
_lock = Lock()


class PlanError(Exception):
    """Понятная пользователю ошибка разбора файла."""


def _norm(value):
    return re.sub(r"\s+", " ", str(value or "")).strip()


def _as_date(value):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = _norm(value)
    if not text:
        return ""
    for pattern in ("%Y-%m-%d", "%d.%m.%Y", "%d/%m/%Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(text[:10], pattern).date().isoformat()
        except ValueError:
            continue
    return ""


def _as_qty(value):
    if value in (None, ""):
        return 0
    if isinstance(value, (int, float)):
        return int(round(value))
    text = _norm(value).replace(" ", "").replace(" ", "").replace(",", ".")
    if not text:
        return 0
    try:
        return int(round(float(text)))
    except ValueError:
        return 0


def _match_columns(header):
    """По строке заголовков находит нужные колонки. Возвращает {роль: индекс}."""
    scores = {}
    for idx, cell in enumerate(header):
        title = _norm(cell).lower()
        if not title:
            continue
        for role, hints in COLUMN_HINTS.items():
            score = sum(weight for hint, weight in hints if hint in title)
            if score > 0:
                scores.setdefault(role, []).append((score, -idx, idx))
    return {role: max(cands)[2] for role, cands in scores.items()}


def parse_workbook(stream, default_date=""):
    """Читает xlsx и возвращает {дата: {"equipment|product": количество}}.

    В рабочем шаблоне «Дата план» — формула =$H$1 c =TODAY(), и если Excel не
    сохранил её вычисленное значение, дата приходит пустой. Тогда берём
    default_date (её задаёт пользователь при загрузке).
    """
    try:
        import openpyxl
    except ImportError as exc:                      # pragma: no cover
        raise PlanError("на сервере нет openpyxl") from exc
    try:
        wb = openpyxl.load_workbook(stream, data_only=True, read_only=True)
    except Exception as exc:
        raise PlanError(f"не удалось открыть файл как Excel: {exc}") from exc

    ws = wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)
    header = next(rows, None)
    if header is None:
        raise PlanError("файл пустой")
    cols = _match_columns(header)
    missing = [r for r in ("product", "quantity") if r not in cols]
    if missing:
        raise PlanError(
            "в заголовке не нашлись колонки: " + ", ".join(missing)
            + ". Нужны «Сокращенное название» и «План количество»")

    fallback = _as_date(default_date) or date.today().isoformat()
    plans, seen, skipped, used_fallback = {}, 0, 0, 0
    for row in rows:
        seen += 1
        if seen > MAX_ROWS:
            raise PlanError(f"слишком много строк (> {MAX_ROWS})")
        def cell(role):
            i = cols.get(role)
            return row[i] if i is not None and i < len(row) else None
        product = _norm(cell("product"))
        qty = _as_qty(cell("quantity"))
        if not product or qty <= 0:
            skipped += 1
            continue
        day = _as_date(cell("date"))
        if not day:
            day = fallback
            used_fallback += 1
        key = f"{_norm(cell('equipment'))}|{product}"
        plans.setdefault(day, {})
        plans[day][key] = plans[day].get(key, 0) + qty
    wb.close()

    if not plans:
        raise PlanError("не нашлось ни одной строки с заполненным планом — "
                        "заполни колонку «План количество»")
    log.info("plan parsed: dates=%s rows_used=%s skipped=%s fallback_date=%s",
             len(plans), sum(len(v) for v in plans.values()), skipped, used_fallback)
    return plans


HISTORY_LIMIT = int(os.getenv("PLAN_HISTORY_LIMIT", "200"))


def _read():
    if not PLAN_FILE.exists():
        return {}
    try:
        with open(PLAN_FILE, encoding="utf-8") as fh:
            data = json.load(fh)
        return data if isinstance(data, dict) else {}
    except Exception:
        log.exception("не удалось прочитать %s", PLAN_FILE)
        return {}


def _write(payload):
    PLAN_FILE.parent.mkdir(parents=True, exist_ok=True)
    # Пишем через временный файл, чтобы не оставить обрезанный JSON при сбое.
    fd, tmp = tempfile.mkstemp(dir=str(PLAN_FILE.parent), suffix=".tmp")
    try:
        with os.fdopen(fd, "w", encoding="utf-8") as fh:
            json.dump(payload, fh, ensure_ascii=False)
        os.replace(tmp, PLAN_FILE)
    except Exception:
        if os.path.exists(tmp):
            os.unlink(tmp)
        raise


def load():
    with _lock:
        return _read().get("plans", {})


def history():
    with _lock:
        entries = _read().get("history", [])
    return list(reversed(entries))          # свежие сверху


def log_upload(entry):
    """Журнал загрузок: и удачных, и отклонённых — чтобы было видно, что пробовали залить."""
    with _lock:
        stored = _read()
        entries = stored.get("history", [])
        entries.append({"at": datetime.now().isoformat(timespec="seconds"), **entry})
        stored["history"] = entries[-HISTORY_LIMIT:]
        stored.setdefault("plans", {})
        _write(stored)


def save(new_plans, source=""):
    """Планы накапливаются: загрузка перезаписывает только присланные даты."""
    with _lock:
        stored = _read()
        current = stored.get("plans", {})
        replaced = {d: sum(current[d].values()) for d in new_plans if d in current}
        current.update(new_plans)
        stored.update({"updated_at": datetime.now().isoformat(timespec="seconds"),
                       "source": source, "plans": current})
        _write(stored)
        return current, replaced


def delete_dates(dates):
    with _lock:
        stored = _read()
        plans = stored.get("plans", {})
        removed = [d for d in dates if plans.pop(d, None) is not None]
        stored.update({"updated_at": datetime.now().isoformat(timespec="seconds"), "plans": plans})
        entries = stored.get("history", [])
        entries.append({"at": datetime.now().isoformat(timespec="seconds"), "action": "delete",
                        "ok": True, "file": "", "dates": removed, "positions": 0})
        stored["history"] = entries[-HISTORY_LIMIT:]
        _write(stored)
        return plans
