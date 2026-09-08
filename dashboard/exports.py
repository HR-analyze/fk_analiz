"""Выгрузки: виджет в Excel и вся сводка в PDF. Только чтение уже посчитанных данных."""
import io
import logging
import os
from datetime import datetime
from zoneinfo import ZoneInfo

log = logging.getLogger("fk.exports")
TZ = ZoneInfo(os.getenv("TIMEZONE", "Europe/Moscow"))
FONT_PATH = os.getenv("REPORT_FONT", "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf")
FONT_BOLD_PATH = os.getenv("REPORT_FONT_BOLD", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf")

SHIFT_NAMES = {"all": "все смены", "day": "дневная (08:00–20:00)", "night": "ночная (20:00–08:00)"}


def _pct(value, total):
    return round(value / total * 100, 1) if total else None


def _done_column(summary):
    """При фильтре, которого нет в плане, столбец процента — не «выполнение»."""
    skew = (summary.get("plan") or {}).get("unsupported_filters") or []
    return ("Доля от полного плана, %" if skew else "Выполнение, %"), skew


def widget_rows(summary, widget):
    """(заголовок, колонки, строки) для одного виджета дашборда."""
    k = summary["kpi"]
    done_column, _ = _done_column(summary)
    if widget == "kpi":
        return "KPI", ["Показатель", "Значение"], [
            ["Операций", k["operations"]],
            ["Из них открытых", k["open_operations"]],
            ["Выпуск, шт", k["quantity"]],
            ["Средняя длительность, мин", k["avg_duration"]],
            ["Средняя производительность, шт/ч", k["avg_rate"]],
            ["Простоев, шт", k["pauses"]],
            ["Простои, мин", k["stop_minutes"]],
            ["Работа, мин", k["work_minutes"]],
            ["Коэффициент использования, %", k["utilization"]],
        ]
    if widget == "daily":
        return "Выпуск по дням", ["Дата", "День", "Операций", "Выпуск, шт", "План, шт",
                                  done_column, "Ср. длительность, мин", "Простои, мин"], [
            [d["date"], d["weekday"], d["ops"], d["qty"], d["plan"] or "", d["done"] if d["done"] is not None else "",
             d["avg_duration"], d["stop"]] for d in summary["daily"]]
    if widget == "hourly":
        target = summary.get("hourly_target", 0)
        return "Производительность по часам", ["Час", "Средняя, шт/ч", "Средняя за период, шт/ч", "Ниже средней"], [
            [h["label"], h["value"], target, "да" if 0 < h["value"] < target else ""] for h in summary["hourly"]]
    if widget in ("equipment", "products", "employees", "reasons"):
        source = {"equipment": ("Выпуск по оборудованию", "Оборудование", "Выпуск, шт", summary["by_equipment"]),
                  "products": ("Топ продуктов", "Продукт", "Выпуск, шт", summary["top_products"]),
                  "employees": ("Выпуск по сотрудникам", "Сотрудник", "Выпуск, шт", summary.get("top_employees", [])),
                  "reasons": ("Простои по причинам", "Причина", "Минуты", summary["pause_reasons"])}[widget]
        title, name_col, value_col, items = source
        total = sum(i["value"] for i in items)
        return title, [name_col, value_col, "Доля, %"], [
            [i["label"], i["value"], _pct(i["value"], total)] for i in items]
    if widget == "work_vs_stop":
        items = summary["work_vs_stop"]
        total = sum(i["value"] for i in items)
        return "Работа и простои", ["Показатель", "Минуты", "Доля, %"], [
            [i["label"], i["value"], _pct(i["value"], total)] for i in items]
    if widget == "employee_timings":
        return "Сводка по сотрудникам", ["Сотрудник", "Операций", "Выпуск, шт", "Доля, %", "Ср. длительность, мин",
                                         "Ср. производительность, шт/ч", "Простоев", "Простои, мин",
                                         "Загрузка, %", "Дней"], [
            [r["employee"], r["ops"], r["qty"], r["share"], r["avg_duration"], r["avg_rate"],
             r["pauses"], r["stop"], r["utilization"], r["days"]] for r in summary.get("employee_timings", [])]
    if widget == "equipment_timings":
        return "Сводка по оборудованию", ["Оборудование", "Операций", "Выпуск, шт", "Работа, мин", "Простои, мин",
                                          "Ср. производительность, шт/ч", "Загрузка, %"], [
            [r["equipment"], r["ops"], r["qty"], r["work"], r["stop"], r["avg_rate"], r["utilization"]]
            for r in summary["equipment_timings"]]
    if widget == "product_timings":
        return "Тайминги по продуктам", ["Продукт", "Операций", "Выпуск, шт", "Ср. длительность, мин",
                                         "Ср. производительность, шт/ч", "Диапазон старта"], [
            [r["product"], r["ops"], r["qty"], r["avg_duration"], r["avg_rate"], r["start_range"]]
            for r in summary["product_timings"]]
    if widget == "plan":
        plan = summary.get("plan") or {}
        return "План и факт", ["Продукт", "Оборудование", "План, шт", "Факт, шт", "Разница, шт", done_column], [
            [r["product"], r["equipment"], r["plan"], r["fact"], r["diff"],
             r["done"] if r["done"] is not None else ""] for r in plan.get("rows", [])]
    if widget == "detail":
        return "Детализация операций", ["Дата", "Оборудование", "Продукт", "Старт", "Финиш",
                                        "Количество, шт", "Статус", "Оператор"], [
            [r["date"], r["equipment"], r["product"], r["start_time"], r["end_time"] or "",
             r["qty"], r["status"], r["user_name"]] for r in summary["detail"]]
    raise KeyError(widget)


ALL_WIDGETS = ("kpi", "daily", "hourly", "equipment", "products", "employees", "reasons",
               "work_vs_stop", "plan", "employee_timings", "equipment_timings",
               "product_timings", "detail")


def filter_rows(filters):
    return [[name, value] for name, value in filters]


def build_xlsx(summary, widget, filters):
    """Один лист с данными виджета и лист с применёнными фильтрами."""
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    widgets = ALL_WIDGETS if widget == "all" else [widget]
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    head_fill = PatternFill("solid", fgColor="E8EEFB")
    head_font = Font(bold=True)

    for name in widgets:
        title, columns, rows = widget_rows(summary, name)
        ws = wb.create_sheet(title[:31])
        ws.append(columns)
        for row in rows:
            ws.append(row)
        for idx, column in enumerate(columns, 1):
            cell = ws.cell(row=1, column=idx)
            cell.fill, cell.font = head_fill, head_font
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            longest = max([len(str(column))] + [len(str(r[idx - 1])) for r in rows[:200]] or [0])
            ws.column_dimensions[get_column_letter(idx)].width = min(max(longest + 2, 11), 46)
        ws.freeze_panes = "A2"
        if rows:
            ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{len(rows) + 1}"

    ws = wb.create_sheet("Фильтры")
    ws.append(["Параметр", "Значение"])
    for row in filter_rows(filters):
        ws.append(row)
    ws.append(["Выгружено", datetime.now(TZ).strftime("%d.%m.%Y %H:%M")])
    _, skew = _done_column(summary)
    if skew:
        ws.append([])
        ws.append(["Внимание", "В файле плана нет разреза по " + ", ".join(skew) + "."])
        ws.append(["", "Факт сужен фильтром, план взят целиком — проценты ниже реальных."])
    for idx in (1, 2):
        ws.cell(row=1, column=idx).fill = head_fill
        ws.cell(row=1, column=idx).font = head_font
        ws.column_dimensions[get_column_letter(idx)].width = 34

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


# ── PDF ───────────────────────────────────────────────────────────────────────

def _register_fonts():
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    regular, bold = "Helvetica", "Helvetica-Bold"
    if os.path.exists(FONT_PATH):
        pdfmetrics.registerFont(TTFont("FK", FONT_PATH))
        regular = "FK"
        if os.path.exists(FONT_BOLD_PATH):
            pdfmetrics.registerFont(TTFont("FK-Bold", FONT_BOLD_PATH))
            bold = "FK-Bold"
        else:
            bold = "FK"
    return regular, bold


def _nice_max(value):
    """Округляем верх оси вверх до «круглого» — иначе подписи вида 41 849."""
    if not value or value <= 0:
        return 1
    import math
    exp = 10 ** math.floor(math.log10(value))
    norm = value / exp
    step = next((s for s in (1, 1.5, 2, 2.5, 3, 4, 5, 6, 8, 10) if norm <= s), 10)
    return step * exp


def _bar_drawing(items, width, height, font, color_hex, target=None, below_hex="#E34948"):
    """Столбцы на примитивах: встроенные графики reportlab плохо держат длинные подписи."""
    from reportlab.graphics.shapes import Drawing, Group, Line, Rect, String
    from reportlab.lib.colors import HexColor

    d = Drawing(width, height)
    items = [i for i in items if i.get("value") is not None]
    if not items or all(not i["value"] for i in items):
        d.add(String(width / 2, height / 2, "Нет данных", fontName=font, fontSize=9,
                     textAnchor="middle", fillColor=HexColor("#667085")))
        return d

    pad_l, pad_r, pad_t, pad_b = 48, 6, 10, 46
    plot_w, plot_h = width - pad_l - pad_r, height - pad_t - pad_b
    top = _nice_max(max([i["value"] for i in items] + ([target] if target else [])))
    slot = plot_w / len(items)
    bar_w = min(slot * 0.62, 26)

    for t in range(5):
        y = pad_b + plot_h * t / 4
        d.add(Line(pad_l, y, width - pad_r, y, strokeColor=HexColor("#EEF1F6"), strokeWidth=0.5))
        d.add(String(pad_l - 4, y - 2.5, f"{top * t / 4:,.0f}".replace(",", " "), fontName=font,
                     fontSize=6, textAnchor="end", fillColor=HexColor("#8A94A6")))

    for i, item in enumerate(items):
        h = item["value"] / top * plot_h
        x = pad_l + slot * i + (slot - bar_w) / 2
        below = bool(target) and 0 < item["value"] < target
        d.add(Rect(x, pad_b, bar_w, max(h, 0.6),
                   fillColor=HexColor(below_hex if below else color_hex), strokeColor=None))
        # Подпись под наклоном — иначе длинные названия налезают друг на друга.
        label = Group(String(0, 0, str(item["label"])[:26], fontName=font, fontSize=5.6,
                             textAnchor="end", fillColor=HexColor("#6B7687")))
        label.translate(pad_l + slot * i + slot / 2 + 2, pad_b - 6)
        label.rotate(40)
        d.add(label)

    d.add(Line(pad_l, pad_b, width - pad_r, pad_b, strokeColor=HexColor("#D9DFE8"), strokeWidth=0.6))
    if target:
        y = pad_b + target / top * plot_h
        d.add(Line(pad_l, y, width - pad_r, y, strokeColor=HexColor("#101828"),
                   strokeWidth=0.8, strokeDashArray=[3, 2]))
    return d


def _pie_drawing(items, width, height, font, colors_hex):
    from reportlab.graphics.shapes import Drawing, Rect, String, Wedge
    from reportlab.lib.colors import HexColor

    d = Drawing(width, height)
    items = [i for i in items if i["value"] > 0]
    if not items:
        d.add(String(width / 2, height / 2, "Нет данных", fontName=font, fontSize=9,
                     textAnchor="middle", fillColor=HexColor("#667085")))
        return d

    total = sum(i["value"] for i in items)
    cx, cy, r = 66, height / 2, min(height / 2 - 8, 52)
    angle = 90.0
    for i, item in enumerate(items):
        sweep = item["value"] / total * 360.0
        d.add(Wedge(cx, cy, r, angle - sweep, angle, fillColor=HexColor(colors_hex[i % len(colors_hex)]),
                    strokeColor=HexColor("#FFFFFF"), strokeWidth=0.8))
        angle -= sweep

    ly = height - 12
    for i, item in enumerate(items[:9]):
        d.add(Rect(width - 190, ly - 3, 6, 6, fillColor=HexColor(colors_hex[i % len(colors_hex)]), strokeColor=None))
        text = f"{str(item['label'])[:30]} — {item['value']:,.0f} ({item['value'] / total * 100:.1f}%)".replace(",", " ")
        d.add(String(width - 180, ly - 2, text, fontName=font, fontSize=6.2, fillColor=HexColor("#3C4658")))
        ly -= 11
    return d


def build_pdf(summary, analysis, filters, title="Сводка по производству"):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import (KeepTogether, PageBreak, Paragraph, SimpleDocTemplate,
                                    Spacer, Table, TableStyle)

    font, bold = _register_fonts()
    out = io.BytesIO()
    doc = SimpleDocTemplate(out, pagesize=A4, leftMargin=13 * mm, rightMargin=13 * mm,
                            topMargin=13 * mm, bottomMargin=13 * mm, title=title)
    width = doc.width

    h1 = ParagraphStyle("h1", fontName=bold, fontSize=17, leading=21, spaceAfter=4)
    h2 = ParagraphStyle("h2", fontName=bold, fontSize=11, leading=14, spaceBefore=8, spaceAfter=4)
    body = ParagraphStyle("body", fontName=font, fontSize=8.5, leading=12, alignment=TA_LEFT)
    muted = ParagraphStyle("muted", fontName=font, fontSize=8, leading=11, textColor=colors.HexColor("#667085"))

    def table(columns, rows, widths=None, limit=40):
        data = [columns] + [[Paragraph(str(c), ParagraphStyle("c", fontName=font, fontSize=6.6, leading=8))
                             for c in row] for row in rows[:limit]]
        t = Table(data, colWidths=widths, repeatRows=1)
        t.setStyle(TableStyle([
            ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#E6EAF0")),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#EFF3FB")),
            ("FONTNAME", (0, 0), (-1, 0), bold),
            ("FONTSIZE", (0, 0), (-1, 0), 6.8),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("TOPPADDING", (0, 0), (-1, -1), 2.5),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 2.5),
        ]))
        return t

    story = [Paragraph(title, h1)]
    applied = ", ".join(f"{n}: {v}" for n, v in filters if v)
    story.append(Paragraph(applied or "фильтры не заданы", muted))
    story.append(Paragraph(f"Сформировано {datetime.now(TZ).strftime('%d.%m.%Y %H:%M')}", muted))
    story.append(Spacer(1, 5 * mm))

    # KPI плиткой 3×2
    k = summary["kpi"]
    tiles = [("Операций", k["operations"]), ("Выпуск, шт", f"{k['quantity']:,}".replace(",", " ")),
             ("Средняя длит., мин", k["avg_duration"]), ("Производительность, шт/ч", k["avg_rate"]),
             ("Простоев", k["pauses"]), ("Коэф. использования, %", k["utilization"])]
    kpi_rows = [[Paragraph(f"<font size=7 color='#667085'>{n}</font><br/><font size=13 name='{bold}'>{v}</font>",
                           ParagraphStyle("k", fontName=font, leading=17)) for n, v in tiles[i:i + 3]]
                for i in (0, 3)]
    kpi = Table(kpi_rows, colWidths=[width / 3] * 3)
    kpi.setStyle(TableStyle([("BOX", (0, 0), (-1, -1), 0.4, colors.HexColor("#E6EAF0")),
                             ("INNERGRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#E6EAF0")),
                             ("TOPPADDING", (0, 0), (-1, -1), 6), ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                             ("LEFTPADDING", (0, 0), (-1, -1), 8)]))
    story += [kpi, Spacer(1, 5 * mm)]

    # Факты
    for heading, lines in analysis:
        story.append(KeepTogether([Paragraph(heading, h2)] +
                                  [Paragraph(line, body) for line in lines]))
    story.append(PageBreak())

    # Графики
    charts = [
        ("Выпуск по дням", [{"label": d["date"][5:], "value": d["qty"]} for d in summary["daily"]],
         "#3B6EF5", None),
        ("Производительность по часам", [{"label": h["label"], "value": h["value"]} for h in summary["hourly"]],
         "#7C5CF5", summary.get("hourly_target") or None),
        ("Выпуск по оборудованию", summary["by_equipment"][:10], "#3FBF94", None),
        ("Топ продуктов", summary["top_products"][:10], "#F0A63C", None),
        ("Выпуск по сотрудникам", summary.get("top_employees", [])[:10], "#2A78D6", None),
    ]
    story.append(Paragraph("Графики", h2))
    for name, items, color_hex, target in charts:
        story.append(KeepTogether([
            Paragraph(name + (f" · средняя {target} шт/ч" if target else ""), body),
            _bar_drawing(items, width, 132, font, color_hex, target),
            Spacer(1, 3 * mm)]))

    pie_palette = ["#2A78D6", "#EB6834", "#1BAF7A", "#EDA100", "#E87BA4", "#008300", "#4A3AA7", "#E34948"]
    story.append(KeepTogether([Paragraph("Простои по причинам", body),
                               _pie_drawing(summary["pause_reasons"], width, 124, font, pie_palette),
                               Spacer(1, 3 * mm)]))
    story.append(KeepTogether([Paragraph("Работа и простои", body),
                               _pie_drawing(summary["work_vs_stop"], width, 110, font, ["#2A78D6", "#E34948"])]))
    story.append(PageBreak())

    # Таблицы
    plan = summary.get("plan") or {}
    tables = [("daily", None), ("employee_timings", None), ("equipment_timings", None), ("product_timings", 25)]
    if plan.get("has_plan"):
        tables.insert(0, ("plan", 30))
    for widget, limit in tables:
        heading, columns, rows = widget_rows(summary, widget)
        if not rows:
            continue
        story.append(Paragraph(heading, h2))
        story.append(table(columns, rows, limit=limit or 40))
        if len(rows) > (limit or 40):
            story.append(Paragraph(f"Показаны первые {limit or 40} из {len(rows)} строк.", muted))
        story.append(Spacer(1, 4 * mm))

    doc.build(story)
    return out.getvalue()
