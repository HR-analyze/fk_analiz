#!/usr/bin/env python3
"""Раздел «Гант»: время операций всех цехов (полночь, окно смены), дорожки параллельных
операций, разобранный чат КЦ, общий тайминг по единому продукту, справочник соответствий,
API, выгрузки и то, что остальные разделы работают как раньше.

    python3 dashboard/tests/test_gantt.py
"""
import asyncio
import io
import json
import os
import sys
import tempfile
from datetime import datetime, timedelta
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent))

_TMP = Path(tempfile.mkdtemp())
_TEST_ENV = {"SOURCE_MODE": "demo", "BHM_SOURCE_MODE": "demo", "PROOFING_SOURCE_MODE": "demo", "KC_SOURCE_MODE": "demo",
             "RELAXDEV_DEPLOY": "", "PLAN_DATABASE_URL": "", "TIMEZONE": "Europe/Moscow", "PLAN_UPLOAD_TOKEN": "tok",
             "PLAN_FILE": str(_TMP / "plans.json"), "EXCLUDE_DATES": ""}
_saved = {k: os.environ.get(k) for k in _TEST_ENV}
os.environ.update(_TEST_ENV)
try:
    import app as dashboard  # noqa: E402
    import gantt  # noqa: E402
    import gantt_exports  # noqa: E402
    import product_map  # noqa: E402
    from sources import TZ  # noqa: E402
finally:
    for _k, _v in _saved.items():
        if _v is None:
            os.environ.pop(_k, None)
        else:
            os.environ[_k] = _v

import bhm_source  # noqa: E402
import kc_source  # noqa: E402
import plans  # noqa: E402
import proofing_source  # noqa: E402
import sources  # noqa: E402

sources.SOURCE_MODE = "demo"
bhm_source.BHM_SOURCE_MODE = "demo"
proofing_source.PROOFING_SOURCE_MODE = "demo"
kc_source.KC_SOURCE_MODE = "demo"
plans.PG_PLANS = False
plans.PLAN_FILE = _TMP / "plans.json"
plans.PLAN_UPLOAD_TOKEN = "tok"
dashboard.RELAXDEV_DEPLOY = False

from aiohttp.test_utils import TestClient, TestServer  # noqa: E402

NOW = datetime(2026, 10, 13, 12, 0, tzinfo=TZ)


def at(day, hh, mm=0):
    return datetime(2026, 10, day, hh, mm, tzinfo=TZ)


def fk_row(product, start, end, created, qty=100, status="closed", equipment="Rondo"):
    return {"id": product + start, "user_name": "Оператор", "equipment": equipment, "product": product,
            "start_time": start, "end_time": end, "qty": qty, "status": status, "created_at": created.isoformat(),
            "date": created.date().isoformat()}


def kc_op(product, operation, start, minutes=None, people=2, pcs=100, status="closed"):
    end = start + timedelta(minutes=minutes) if minutes else None
    return {"id": f"{product}-{operation}-{start:%H%M}", "product": product, "operation": operation, "people": people,
            "pcs": pcs if status == "closed" else None, "status": status, "start_at": start, "end_at": end,
            "closed_at": end, "duration_min": minutes, "valid": bool(minutes), "person_hours": None,
            "user_name": "Анна", "closed_by_name": "Мария" if end else ""}


# ── время и источники ─────────────────────────────────────────────────────────

def test_fk_clock_handles_midnight_and_record_time():
    # запись сделана в конце операции (00:15 следующих суток) — начало в прошлые сутки
    s, e = gantt._clock(at(8, 0, 15), "23:00", "00:15")
    assert s == at(7, 23) and e == at(8, 0, 15)
    # запись сделана в начале операции
    s, e = gantt._clock(at(7, 23, 0), "23:00", "00:15")
    assert s == at(7, 23) and e == at(8, 0, 15)
    assert gantt._clock(at(7, 9), "", "")[0] is None


def test_bars_from_every_section_people_only_in_kc():
    fk = {"production": [fk_row("Багет", "09:00", "10:30", at(7, 10, 30), qty=900)]}
    bhm = {"production": [{"id": "b1", "user_name": "Ольга", "equipment": "Diosna", "product": "Багет", "operation": "Замес теста",
                           "start_time": "07:00", "end_time": "07:20", "qty": 150.5, "status": "closed",
                           "created_at": at(7, 7).isoformat(), "date": "2026-10-07"}]}
    kc = {"operations": [kc_op("Эклер", "Начинка", at(7, 11), 60, people=3, pcs=600)]}
    bars = gantt.collect_bars({"fk": fk, "bhm": bhm, "kc": kc}, {"operations": [], "handoffs": []})
    f, b, k = bars["fk"][0], bars["bhm"][0], bars["kc"][0]
    assert (f["operation"], f["qty"], f["unit"], f["people"], f["equipment"]) == ("Формовка", 900, "шт", None, "Rondo")
    assert (b["operation"], b["unit"], b["equipment"]) == ("Замес теста", "кг", "Diosna")
    assert (k["people"], k["qty"]) == (3, 600)
    v = gantt.view(k, NOW)
    assert v["rate"] == 600.0 and v["rate_per_person"] == 200.0 and v["person_hours"] == 3.0
    assert gantt.view(f, NOW)["rate_per_person"] is None   # людей у ФК нет — на человека не считаем


def test_proofing_stages_become_two_operations():
    p = json.loads((HERE.parent / "proofing_demo_data.json").read_text(encoding="utf-8"))
    snap = proofing_source.normalize(p["operations"], p["batches"], p["pauses"], p["photos"], p["catalog"])
    bars = gantt.proofing_bars(snap)
    done = next(b for b in snap["batches"] if b["status"] == "done")
    mine = [x for x in bars if x["id"].startswith(f"pr-{done['id']}")]
    assert {x["stage"] for x in mine} == {"proofing", "baking"}
    baking = next(x for x in mine if x["stage"] == "baking")
    assert baking["qty"] == done["pieces_out"] and baking["equipment"].endswith("печь")
    assert not any(x["id"].startswith(f"pr-{b['id']}") for b in snap["batches"] if b["status"] == "cancelled" for x in bars)


# ── чат КЦ ────────────────────────────────────────────────────────────────────

def test_chat_history_is_consistent():
    h = gantt.load_history()
    ops = h["operations"]
    assert len(ops) >= 170 and len({o["id"] for o in ops}) == len(ops)
    for o in ops:
        assert o["flag"] in ("ok", "check", "invalid") and o["msgs"], o
        if o["start_at"] and o["end_at"]:
            assert o["end_at"] >= o["start_at"], o
        assert (o["status"] == "closed") == bool(o["end_at"]), o
        assert o["people"] is None or 1 <= o["people"] <= 50, o
        assert o["unit"] in ("", "шт", "кг", "пласт", "лист", "замес", "форма", "кольцо", "шпилька"), o
    assert {h_["product"] for h_ in h["handoffs"]} >= {"Чизкейк Сан-Себастьян", "Трубочка с заварным кремом"}
    # тексты сообщений и имена в файл не попадают
    raw = (HERE.parent / "kc_chat_history.json").read_text(encoding="utf-8")
    assert "Константин" not in raw and "text" not in json.loads(raw)["operations"][0]


def test_chat_bars_invalid_and_approx():
    h = gantt.load_history()
    bars = gantt.kc_bars(None, h)
    invalid = [b for b in bars if b["flag"] == "invalid"]
    assert invalid and all(gantt.view(b, NOW)["rate"] is None for b in invalid)
    approx = [b for b in bars if b["approx"]]
    assert approx and all(b["flag"] == "check" for b in approx)
    no_end = [b for b in bars if not b["end"]]
    assert no_end and all(not b["running"] and gantt.bar_end(b, NOW) is None for b in no_end)


# ── окно, дорожки ─────────────────────────────────────────────────────────────

def test_window_full_day_night_and_open_ops():
    kc = {"operations": [kc_op("Эклер", "Начинка", at(7, 23), 120),          # 23:00–01:00
                         kc_op("Эклер", "Глазировка", at(7, 10), 60),
                         kc_op("Эклер", "Декор", at(7, 15), status="open")]}  # бот: идёт до «сейчас»
    bars = gantt.collect_bars({"kc": kc}, {"operations": [], "handoffs": []})

    def ops(mode):
        s = gantt.build(bars, {}, {}, "2026-10-07", mode, ["kc"], now=NOW)
        return sorted(b["operation"] for r in s["rows"] for b in r["bars"])
    assert ops("full") == ["Глазировка", "Декор", "Начинка"]
    assert ops("day") == ["Глазировка", "Декор"]
    assert ops("night") == ["Декор", "Начинка"]
    s = gantt.build(bars, {}, {}, "2026-10-07", "full", ["kc"], now=NOW)
    running = next(b for r in s["rows"] for b in r["bars"] if b["operation"] == "Декор")
    assert running["running"] and running["draw_end"] == NOW.isoformat()


def test_parallel_operations_get_separate_lanes():
    kc = {"operations": [kc_op("Ром баба", "Накалывание", at(7, 9), 120), kc_op("Ром баба", "Накалывание", at(7, 10), 60),
                         kc_op("Ром баба", "Накалывание", at(7, 12), 30)]}
    s = gantt.build(gantt.collect_bars({"kc": kc}, {}), {}, {}, "2026-10-07", "full", ["kc"], now=NOW)
    row = s["rows"][0]
    assert row["lanes"] == 2 and [b["lane"] for b in row["bars"]] == [0, 1, 0]


# ── строки: продукт → операции, колонки ───────────────────────────────────────

def test_rows_grouped_by_product_with_columns():
    kc = {"operations": [kc_op("Булка Много мака", "Отделка глазурью", at(7, 9), 60, people=2, pcs=100),
                         kc_op("Булка Много мака", "Отделка соусом", at(7, 9, 30), 30, people=3, pcs=90)]}
    fk = {"production": [fk_row("Багет", "09:00", "10:00", at(7, 10), qty=900)]}
    s = gantt.build(gantt.collect_bars({"kc": kc, "fk": fk}, {}), {}, {}, "2026-10-07", "full", now=NOW)
    rows = {(r["product"], r["operation"]): r for r in s["rows"]}
    glaze, sauce = rows[("Булка Много мака", "Отделка глазурью")], rows[("Булка Много мака", "Отделка соусом")]
    assert (glaze["people"], glaze["rate_per_person"], glaze["rate_unit"]) == (2, 50.0, "шт")    # 100 шт / (1 ч × 2)
    assert (sauce["people"], sauce["rate_per_person"]) == (3, 60.0)                              # 90 / (0,5 ч × 3)
    prod = next(x for x in s["products"] if x["product"] == "Булка Много мака")
    assert prod["ops"] == 2 and prod["section_title"] == "Кондитерский цех"
    assert prod["people"] == 5 and prod["people_ops"] == 5          # 09:30–10:00 работают обе операции
    # выработка продукта — самая массовая операция (глазурь, 100 шт), операции не складываются:
    # 100 шт / (2 + 1,5) чел·ч всех операций; в час — / 60 мин работы с продуктом
    assert (prod["qty"], prod["rate_per_person"], prod["rate"]) == ("100 шт", 28.6, 100.0)
    assert prod["output_from"] == "Кондитерский цех · Отделка глазурью"
    bread = next(x for x in s["products"] if x["product"] == "Багет")
    assert bread["section_title"] == "Формовочный цех" and bread["people"] is None
    assert bread["rate_per_person"] is None and bread["rate"] == 900.0 and bread["equipment"] == "Rondo"

    # тестозаготовочный цех: людей нет — кг продукта в час работы с ним (замес 200 кг, работа 08:00–09:00)
    bhm = {"production": [{"id": f"b{i}", "user_name": "О", "equipment": eq, "product": "Багет", "operation": op,
                           "start_time": st, "end_time": en, "qty": qty, "status": "closed",
                           "created_at": at(7, 9).isoformat(), "date": "2026-10-07"}
                          for i, (op, eq, st, en, qty) in enumerate([("Сборка теста", "Daxner", "08:00", "08:20", None),
                                                                     ("Замес теста", "Diosna", "08:20", "08:40", 200.0),
                                                                     ("Отлежка теста", "Без оборудования", "08:40", "09:00", None)])]}
    s = gantt.build(gantt.collect_bars({"bhm": bhm}, {}), {}, {}, "2026-10-07", "full", now=NOW)
    dough = s["products"][0]
    assert (dough["ops"], dough["qty"], dough["rate"], dough["rate_per_person"]) == (3, "200 кг", 200.0, None)

    # незакрытая с утра «Сборка теста» идёт до «сейчас» — в выработку в час её время не входит
    forgot = {"production": bhm["production"] + [
        {"id": "b9", "user_name": "О", "equipment": "Daxner", "product": "Багет", "operation": "Сборка теста",
         "start_time": "08:00", "end_time": "", "qty": None, "status": "open", "created_at": at(7, 8).isoformat(), "date": "2026-10-07"}]}
    s = gantt.build(gantt.collect_bars({"bhm": forgot}, {}), {}, {}, "2026-10-07", "day", now=at(7, 18))
    dough = s["products"][0]
    assert dough["busy_min"] == 600 and dough["running"] == 1          # на шкале — 08:00–18:00
    assert (dough["rate"], dough["work_min"]) == (200.0, 60)           # 200 кг за 1 ч закрытых операций


def test_peak_people_relay_is_not_parallel():
    kc = {"operations": [kc_op("Эклер", "Начинка", at(7, 9), 60, people=2), kc_op("Эклер", "Декор", at(7, 10), 60, people=3)]}
    s = gantt.build(gantt.collect_bars({"kc": kc}, {}), {}, {}, "2026-10-07", "full", now=NOW)
    assert s["products"][0]["people"] == 3 and s["products"][0]["people_ops"] == 5


def test_busy_time_is_clipped_to_window():
    kc = {"operations": [kc_op("Эклер", "Декор", at(5, 9), status="open")]}   # идёт третьи сутки
    s = gantt.build(gantt.collect_bars({"kc": kc}, {}), {}, {}, "2026-10-07", "day", now=NOW)
    assert s["rows"][0]["busy_min"] == 12 * 60


def test_hand_typed_names_are_merged_and_ops_mapped():
    kc = {"operations": [kc_op("Булка Много мака", "Начинка", at(7, 9), 30, pcs=50),
                         kc_op("булка  много-мака", "начинка", at(7, 10), 30, pcs=60),
                         kc_op("Булка Много мака", "Внесение начинки", at(7, 11), 30, pcs=70)]}
    bars = gantt.collect_bars({"kc": kc}, {})
    s = gantt.build(bars, {}, {}, "2026-10-07", "full", now=NOW)
    assert {r["product"] for r in s["rows"]} == {"Булка Много мака"}
    assert sorted(r["operation"] for r in s["rows"]) == ["Внесение начинки", "Начинка"]
    s = gantt.build(bars, {}, {}, "2026-10-07", "full", now=NOW, ops_map={"kc": {"начинка": "Внесение начинки"}})
    assert [(r["operation"], r["runs"]) for r in s["rows"]] == [("Внесение начинки", 3)]
    raw = {b["operation_raw"] for r in s["rows"] for b in r["bars"]}
    assert raw == {"Начинка", "начинка", "Внесение начинки"}       # исходное название видно в подсказке
    m = gantt.meta(bars, {}, {}, None, {})
    assert m["products"] == ["Булка Много мака"]


# ── среднее по продукту ───────────────────────────────────────────────────────

def test_averages_per_product_relative_to_start():
    kc = {"operations": [kc_op("Торт", "Сборка", at(7, 9), 60, people=2, pcs=20),
                         kc_op("Торт", "Декор", at(7, 10, 30), 30, people=1, pcs=20),
                         kc_op("Торт", "Сборка", at(8, 10), 90, people=2, pcs=30),
                         kc_op("Торт", "Декор", at(8, 12), 20, people=1, pcs=30),
                         kc_op("Торт", "Сборка", at(11, 9), 60, people=2, pcs=20)]}   # вне периода
    bars = gantt.collect_bars({"kc": kc}, {})
    av = gantt.averages(bars, gantt.Namer(bars), "2026-10-07", "2026-10-08", now=NOW)
    assert av["days"] == 2 and len(av["products"]) == 1
    p = av["products"][0]
    assert p["days"] == 2 and p["lead_min"] == 130.0 and (p["lead_min_min"], p["lead_min_max"]) == (120, 140)
    ops = {o["operation"]: o for o in p["ops"]}
    assert ops["Сборка"]["offset_min"] == 0 and ops["Сборка"]["span_min"] == 75.0
    assert ops["Декор"]["offset_min"] == 105.0 and ops["Декор"]["span_min"] == 25.0   # (90 + 120) / 2
    assert ops["Сборка"]["rate_per_person"] == round(50 / 5, 1)                        # Σшт / Σчел·ч за дни
    assert ops["Сборка"]["qty_day"] == "25 шт"
    assert p["stages"][0]["key"] == "kc" and p["stages"][0]["span_min"] == 130.0
    # выработка продукта: самая массовая операция последнего этапа, операции не складываются
    o = p["output"]
    assert (o["qty_day"], o["unit"], o["operation"], o["days"]) == (25.0, "шт", "Декор", 2)   # поровну — позже закончилась
    assert o["rate_per_person"] == round(50 / (2 * 1 + 0.5 * 1 + 3 * 1 + 20 / 60), 1)   # шт / чел·ч всех операций
    s = gantt.build(bars, {}, {}, "2026-10-07", "full", now=NOW, avg_from="2026-10-08", avg_to="2026-10-07")
    assert (s["averages"]["from"], s["averages"]["to"]) == ("2026-10-07", "2026-10-08")


def test_product_output_takes_last_stage_and_skips_batches():
    fk = gantt.fk_bars({"production": [fk_row("Батон", "09:00", "10:00", at(7, 10), qty=1000)]})
    baking = [gantt._bar("proofing", "baking", "Батон", "Отпекание", at(7, 12), at(7, 13), qty=950, unit="шт")]
    batter = [gantt._bar("kc", "kc", "Бисквит", "Разлив", at(7, 9), at(7, 10), qty=4, unit="замес", people=2)]
    bars = {"fk": fk, "proofing": baking, "kc": batter}
    av = {p["product"]: p for p in gantt.averages(bars, gantt.Namer(bars), "2026-10-07", "2026-10-07", now=NOW)["products"]}
    o = av["Батон"]["output"]
    assert (o["qty_day"], o["stage"], o["rate_per_person"]) == (950.0, "baking", None)   # людей нет — на чел·ч не считаем
    assert av["Бисквит"]["output"] is None                                              # замес — не продукция


# ── общий тайминг ─────────────────────────────────────────────────────────────

def _chain():
    bhm = {"production": [{"id": "b1", "user_name": "О", "equipment": "Diosna", "product": "Батон / ромашка",
                           "operation": "Замес теста", "start_time": "08:00", "end_time": "08:30", "qty": 200.0,
                           "status": "closed", "created_at": at(7, 8).isoformat(), "date": "2026-10-07"}]}
    fk = {"production": [fk_row("Батон классический", "09:00", "10:00", at(7, 10), qty=1000)]}
    return gantt.collect_bars({"bhm": bhm, "fk": fk}, {"operations": [], "handoffs": []})


def test_timing_links_sections_only_through_mapping():
    bars = _chain()
    loose = gantt.build(bars, {}, {}, "2026-10-07", "full", now=NOW)["timing"]
    assert {t["product"] for t in loose} == {"Батон / ромашка", "Батон классический"}
    assert all(t["stages_count"] == 1 for t in loose)
    mapping = {"bhm": {"Батон / ромашка": "Батон классический"}}
    linked = gantt.build(bars, {}, mapping, "2026-10-07", "full", now=NOW)["timing"]
    assert len(linked) == 1
    t = linked[0]
    st = {s["key"]: s for s in t["stages"]}
    assert t["stages_count"] == 2 and t["lead_min"] == 120             # 08:00 → 10:00
    assert st["fk"]["gap_min"] == 30                                   # ожидание между тестом и формовкой
    assert st["bhm"]["qty"] == "200 кг" and st["fk"]["qty"] == "1 000 шт"   # единицы не складываются
    assert st["pack"]["present"] is False and st["pack"]["source"] == "нет источника"


def test_packing_handoffs_from_chat():
    h = {"operations": [], "handoffs": [{"product": "Трубочка", "at": at(11, 14, 54).isoformat(), "qty": 930, "unit": "шт",
                                         "exact": False}]}
    kc = {"operations": [kc_op("Трубочка", "Начинка", at(11, 10, 39), 240, people=1, pcs=930)]}
    s = gantt.build(gantt.collect_bars({"kc": kc}, h), h, {}, "2026-10-11", "full", now=NOW)
    t = s["timing"][0]
    pack = next(x for x in t["stages"] if x["key"] == "pack")
    assert pack["handoffs"][0]["qty"] == 930 and s["handoffs"][0]["product"] == "Трубочка"


# ── справочник ────────────────────────────────────────────────────────────────

def test_product_map_template_roundtrip_and_errors():
    import openpyxl
    blob = product_map.template_xlsx({"bhm": ["Батон / ромашка"], "fk": ["Батон классический"]}, {})
    wb = openpyxl.load_workbook(io.BytesIO(blob))
    ws = wb.active
    assert [c.value for c in ws[1]][:3] == ["Цех", "Продукт в цехе", "Единый продукт"]
    ws["C2"] = "Батон классический"
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    mapping, count = product_map.parse_xlsx(out)
    assert mapping == {"bhm": {"Батон / ромашка": "Батон классический"}} and count == 1
    bad = openpyxl.Workbook()
    bad.active.append(["Цех", "Продукт в цехе", "Единый продукт"])
    bad.active.append(["Склад", "Х", "Y"])
    buf = io.BytesIO()
    bad.save(buf)
    buf.seek(0)
    try:
        product_map.parse_xlsx(buf)
        raise AssertionError("ожидалась ошибка цеха")
    except product_map.MapError as exc:
        assert "Склад".lower() in str(exc)


def test_map_operations_sheet_hints_and_old_section_name():
    import openpyxl
    blob = product_map.template_xlsx({"kc": ["Круассан с малиной", "Круассан с малиной фреш"]}, {},
                                     {"kc": {"Начинка": ["Эклер"], "Внесение начинки": ["Эклер"]}}, {})
    wb = openpyxl.load_workbook(io.BytesIO(blob))
    assert wb.sheetnames[:2] == ["Соответствия", "Операции"]
    hints = {r[1]: r[3] for r in wb["Соответствия"].iter_rows(min_row=2, values_only=True)}
    assert hints["Круассан с малиной"] == "Круассан с малиной фреш"    # подсказка, не склейка
    ops = wb["Операции"]
    for row in ops.iter_rows(min_row=2):
        if row[1].value == "Начинка":
            row[2].value = "Внесение начинки"
    wb["Соответствия"].append(["ФК", "Батон", "Батон классический"])    # старые названия цехов
    wb["Соответствия"].append(["Формовка", "Багет", "Багет французский"])
    wb["Соответствия"].append(["Расстойка", "Хала", "Хала плетёная"])
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    parsed = product_map.parse_workbook(out)
    assert parsed["ops"] == {"kc": {"Начинка": "Внесение начинки"}} and parsed["op_pairs"] == 1
    assert parsed["map"] == {"fk": {"Батон": "Батон классический", "Багет": "Багет французский"},
                             "proofing": {"Хала": "Хала плетёная"}}
    assert product_map.unify({"kc": {"Начинка": "X"}}, "kc", "  начинка ") == "X"
    assert product_map.norm("Булка «Много-мака»  ё") == "булка много мака е"


def test_product_map_file_storage_does_not_touch_plans():
    plans_file = Path(plans.PLAN_FILE)
    before = plans_file.read_text(encoding="utf-8") if plans_file.exists() else None
    product_map.save({"fk": {"А": "Б"}}, "x.xlsx")
    assert product_map.load()["map"] == {"fk": {"А": "Б"}}
    after = plans_file.read_text(encoding="utf-8") if plans_file.exists() else None
    assert before == after
    product_map.save({}, "")


# ── API и выгрузки ────────────────────────────────────────────────────────────

async def _client():
    app = dashboard.build_app()
    app.on_startup.clear()
    client = TestClient(TestServer(app))
    await client.start_server()
    return client


class patched:
    def __init__(self, module, **values):
        self.module, self.values, self.old = module, values, {}

    def __enter__(self):
        for k, v in self.values.items():
            self.old[k] = getattr(self.module, k)
            setattr(self.module, k, v)

    def __exit__(self, *exc):
        for k, v in self.old.items():
            setattr(self.module, k, v)


def test_api_meta_summary_and_one_section_down():
    async def run():
        dashboard.kc_source = kc_source.KcSource()
        client = await _client()
        try:
            m = await (await client.get("/api/meta?source=gantt")).json()
            assert m["ok"] and m["kind"] == "gantt" and m["days"] and m["history_ops"] >= 170
            d = await (await client.get("/api/summary?source=gantt&date=2026-10-07&window=full")).json()
            assert d["ok"] and d["totals"]["ops"] > 0 and d["totals"]["from_chat"] > 0
            assert set(d["totals"]["by_section"]) == {"Тестозаготовочный цех", "Формовочный цех", "Печной цех", "Кондитерский цех"}
            assert d["products"] and d["averages"]["products"] and m["avg_from"] <= m["avg_to"]
            av = await (await client.get("/api/summary?source=gantt&date=2026-10-07&avg_from=2026-10-07&avg_to=2026-10-07")).json()
            assert (av["averages"]["from"], av["averages"]["to"], av["averages"]["days"]) == ("2026-10-07", "2026-10-07", 1)
            # данные КЦ из переписки — неточные, и это написано на странице и в выгрузках
            assert "переписки" in d["chat_note"] and "неточные" in d["chat_note"] and "переписки" in m["chat_note"]
            no_kc = await (await client.get("/api/summary?source=gantt&date=2026-10-07&sections=bhm,fk")).json()
            assert no_kc["chat_note"] == ""
            assert any(o["from_chat"] for p in d["averages"]["products"] for o in p["ops"])
            only = await (await client.get("/api/summary?source=gantt&date=2026-10-07&sections=kc")).json()
            assert {r["section"] for r in only["rows"]} == {"kc"}
            # КЦ упал — Гант работает по остальным цехам и честно пишет, чего нет
            with patched(kc_source, KC_SOURCE_MODE="api", KC_API_URL="http://127.0.0.1:9", KC_API_KEY="k", KC_TIMEOUT=3):
                dashboard.kc_source = kc_source.KcSource()
                down = await (await client.get("/api/summary?source=gantt&date=2026-10-07")).json()
                assert down["ok"] and "kc" in down["section_errors"]
                assert {r["section"] for r in down["rows"]} >= {"proofing"}
                assert all(b["source"] == "chat" for r in down["rows"] if r["section"] == "kc" for b in r["bars"])
        finally:
            dashboard.kc_source = kc_source.KcSource()
            await client.close()
    asyncio.run(run())


def test_api_map_upload_needs_token_and_changes_timing():
    async def run():
        import openpyxl
        client = await _client()
        try:
            r = await client.get("/api/gantt/map/template")
            wb = openpyxl.load_workbook(io.BytesIO(await r.read()))
            ws = wb.active
            for row in ws.iter_rows(min_row=2):
                if row[0].value == "Тестозаготовочный цех" and row[1].value == "Круассан классический":
                    row[2].value = "Круассан французский"
            op_ws = wb["Операции"]
            for row in op_ws.iter_rows(min_row=2):
                if row[0].value == "Кондитерский цех" and row[1].value == "Начинка":
                    row[2].value = "Внесение начинки"
            buf = io.BytesIO()
            wb.save(buf)
            data = buf.getvalue()

            async def upload(token):
                from aiohttp import FormData
                form = FormData()
                form.add_field("file", data, filename="map.xlsx")
                return await client.post("/api/gantt/map/upload", data=form, headers={"X-Plan-Token": token})
            before = await (await client.get("/api/summary?source=gantt&date=2026-10-09&sections=kc")).json()
            assert any(r["operation"] == "Начинка" for r in before["rows"])
            assert (await upload("wrong")).status == 401
            r = await upload("tok")
            res = await r.json()
            assert r.status == 200 and res["pairs"] == 1 and res["op_pairs"] == 1
            state = await (await client.get("/api/gantt/map")).json()
            assert state["map"] == {"bhm": {"Круассан классический": "Круассан французский"}}
            assert state["ops"] == {"kc": {"Начинка": "Внесение начинки"}} and "Начинка" in state["operations"]["kc"]
            d = await (await client.get("/api/summary?source=gantt&date=2026-10-09&sections=kc")).json()
            assert not any(r["operation"] == "Начинка" for r in d["rows"])
            assert any(r["operation"] == "Внесение начинки" for r in d["rows"])
            meta = await (await client.get("/api/meta?source=gantt")).json()
            assert "Круассан классический" not in meta["products"]
        finally:
            product_map.save({}, "")
            await client.close()
    asyncio.run(run())


def test_exports_xlsx_and_pdf():
    async def run():
        import openpyxl
        client = await _client()
        try:
            r = await client.get("/api/export/xlsx?source=gantt&date=2026-10-07&window=full")
            assert r.status == 200 and "gantt_2026-10-07_full" in r.headers["Content-Disposition"]
            wb = openpyxl.load_workbook(io.BytesIO(await r.read()))
            assert wb.sheetnames == ["Процессы по цехам", "Цеха · операции", "Продукты (среднее)", "Этапы по продуктам",
                                     "Строки диаграммы", "Операции", "Общий тайминг", "Передано на фасовку",
                                     "Справочник продуктов", "Справочник операций", "Фильтры"]
            page = await (await client.get("/api/summary?source=gantt&date=2026-10-07&window=full")).json()
            assert wb["Операции"].max_row - 1 == page["totals"]["ops"]
            assert wb["Общий тайминг"].max_row - 1 == len(page["timing"])
            assert wb["Строки диаграммы"].max_row - 1 == len(page["rows"]) + len(page["products"])
            av = page["averages"]["products"]
            # плоские листы процессов: строка — продукт × операция; время — числом (ч:мм)
            proc = wb["Процессы по цехам"]
            assert proc.max_row - 1 == sum(len(x["ops"]) for x in av)
            assert [c.value for c in proc[1]][:4] == ["Цех", "Этап цеха", "Продукт", "Операция"]
            assert proc.cell(row=2, column=9).number_format == "[h]:mm"
            shops = {r[0] for r in proc.iter_rows(min_row=2, values_only=True)}
            assert shops <= {"Тестозаготовочный цех", "Формовочный цех", "Печной цех", "Кондитерский цех"} and len(shops) >= 3
            assert wb["Продукты (среднее)"].max_row - 1 == len(av)
            assert wb["Этапы по продуктам"].max_row - 1 == sum(len(x["stages"]) for x in av)
            assert wb["Цеха · операции"].max_row > 1
            notes = [c.value for c in wb["Фильтры"]["B"] if c.value]
            assert any("переписки" in str(v) for v in notes)
            assert any(r[-1] == "переписка кондитерского цеха — неточно" for r in wb["Процессы по цехам"].iter_rows(min_row=2, values_only=True))
            r = await client.get("/api/export/pdf?source=gantt&date=2026-10-07")
            blob = await r.read()
            assert r.status == 200 and blob[:4] == b"%PDF"
            import pypdf
            text = "\n".join(pg.extract_text() for pg in pypdf.PdfReader(io.BytesIO(blob)).pages)
            assert "Процессы по продуктам" in text and "Этап · операция" in text
            for x in page["averages"]["products"][:5]:   # блок на каждый продукт
                assert x["product"] in text, x["product"]
        finally:
            await client.close()
    asyncio.run(run())


def test_other_sections_unchanged():
    async def run():
        client = await _client()
        try:
            for src in ("fk", "bhm", "proofing", "kc"):
                r = await client.get(f"/api/summary?source={src}")
                assert r.status == 200 and (await r.json())["ok"], src
            assert (await (await client.get("/api/summary")).json())["source_key"] == "fk"
            meta = await (await client.get("/api/meta?source=fk")).json()
            assert meta["sources"]["gantt"] is True
        finally:
            await client.close()
    asyncio.run(run())


if __name__ == "__main__":
    tests = [v for k, v in globals().items() if k.startswith("test_")]
    for t in tests:
        t()
        print(f"  ✅ {t.__name__}")
    print(f"ИТОГ: {len(tests)} тестов прошли")
