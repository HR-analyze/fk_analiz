#!/usr/bin/env python3
"""Раздел «КЦ» (кондитерский цех): API и ключ, статусы операций, человеко-часы и выработка,
отсутствие ложного «готового выпуска», нормы «продукция + операция», смены и полночь,
период по времени завершения, сотрудники, остановки, Excel/PDF, отказ API и то, что
ФК, БХМ и расстойка работают как раньше.

Без базы и без сети наружу (поддельный бот поднимается локально):
    python3 dashboard/tests/test_kc.py
"""
import asyncio
import io
import json
import os
import sys
import tempfile
from datetime import datetime, timedelta, timezone
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent))

_TEST_ENV = {"SOURCE_MODE": "demo", "BHM_SOURCE_MODE": "demo", "PROOFING_SOURCE_MODE": "demo",
             "KC_SOURCE_MODE": "api", "KC_API_URL": "http://127.0.0.1:9", "KC_API_KEY": "secret-kc-key",
             "RELAXDEV_DEPLOY": "", "PLAN_DATABASE_URL": "", "TIMEZONE": "Europe/Moscow",
             "PLAN_FILE": str(Path(tempfile.mkdtemp()) / "plans.json"), "EXCLUDE_DATES": ""}
_saved = {k: os.environ.get(k) for k in _TEST_ENV}
os.environ.update(_TEST_ENV)
try:
    import app as dashboard  # noqa: E402
    import kc_aggregate as ka  # noqa: E402
    import kc_exports as ke  # noqa: E402
    import kc_source as ks  # noqa: E402
    from sources import TZ  # noqa: E402
finally:
    for _k, _v in _saved.items():
        if _v is None:
            os.environ.pop(_k, None)
        else:
            os.environ[_k] = _v

# Модули могли быть импортированы раньше другим тестом: режимы задаём явно.
import bhm_source  # noqa: E402
import plans  # noqa: E402
import proofing_source  # noqa: E402
import sources  # noqa: E402

sources.SOURCE_MODE = "demo"
bhm_source.BHM_SOURCE_MODE = "demo"
proofing_source.PROOFING_SOURCE_MODE = "demo"
plans.PG_PLANS = False
plans.PLAN_FILE = Path(_TEST_ENV["PLAN_FILE"])
dashboard.RELAXDEV_DEPLOY = False

from aiohttp import web  # noqa: E402
from aiohttp.test_utils import TestClient, TestServer  # noqa: E402

UTC = timezone.utc
KEY = "secret-kc-key"
NOW = datetime(2026, 10, 20, 12, 0, tzinfo=TZ)
# Поля операции и остановки — ровно как в test_dashboard_all_contract бота кондитерки.
DASH_OPERATION_FIELDS = {"id", "product", "operation", "people", "quantity_pcs", "start_time", "end_time", "start_at",
                         "end_at", "status", "user_id", "user_name", "created_at", "closed_at", "closed_by_id",
                         "closed_by_name", "duration_min", "person_hours", "pcs_per_hour", "pcs_per_person_hour"}
DASH_PAUSE_FIELDS = {"id", "user_id", "user_name", "reason", "start_time", "end_time", "start_at", "end_at", "created_at"}


def local(day, hh, mm=0, month=10):
    return datetime(2026, month, day, hh, mm, tzinfo=TZ)


def iso(dt):
    return dt.astimezone(UTC).isoformat() if dt else None


_n = [0]


def op(product, operation, start, minutes=None, pcs=None, people=1, status="closed", user="Анна", closer="Мария"):
    """Строка API бота: closed — с концом и штуками, open — без, cancelled — с closed_at."""
    _n[0] += 1
    end = start + timedelta(minutes=minutes) if status == "closed" else None
    closed_at = end if status == "closed" else (start + timedelta(minutes=5) if status == "cancelled" else None)
    row = {"id": f"op-{_n[0]}", "product": product, "operation": operation, "people": people,
           "quantity_pcs": pcs if status == "closed" else None,
           "start_time": start.strftime("%H:%M"), "end_time": end.strftime("%H:%M") if end else None,
           "start_at": iso(start), "end_at": iso(end), "status": status, "user_id": 1, "user_name": user,
           "created_at": iso(start), "closed_at": iso(closed_at), "closed_by_id": 2 if closed_at else None,
           "closed_by_name": closer if closed_at else None}
    # посчитанные ботом поля (bot.metrics) — дашборд их не использует, но они есть в ответе
    if status == "closed":
        hours = minutes / 60
        ph = round(hours * people, 2)
        row.update(duration_min=minutes, person_hours=ph, pcs_per_hour=round(pcs / hours, 1) if pcs else None,
                   pcs_per_person_hour=round(pcs / ph, 1) if pcs else None)
    else:
        row.update(duration_min=None, person_hours=None, pcs_per_hour=None, pcs_per_person_hour=None)
    return row


def pause(start, minutes, reason="Нет сырья", user="Анна"):
    _n[0] += 1
    end = start + timedelta(minutes=minutes)
    return {"id": f"p-{_n[0]}", "user_id": 1, "user_name": user, "reason": reason, "start_time": start.strftime("%H:%M"),
            "end_time": end.strftime("%H:%M"), "start_at": iso(start), "end_at": iso(end), "created_at": iso(end)}


def snap(ops, pauses=()):
    return ks.normalize(ops, list(pauses), {"reasons": ["Нет сырья"]})


def summary(snapshot, **filters):
    filters.setdefault("date_from", "2026-10-01")
    filters.setdefault("date_to", "2026-10-15")
    return ka.build_summary(snapshot, filters, now=NOW)


def rum_baba(day=5):
    """Пример цеха: одна партия ромовой бабы проходит три операции по 600 шт."""
    return [op("Ромовая баба", "Накалывание", local(day, 8), 180, 600, 1),
            op("Ромовая баба", "Купание", local(day, 11, 30), 120, 600, 2),
            op("Ромовая баба", "Декорирование", local(day, 14), 60, 600, 3)]


# ── контракт API ──────────────────────────────────────────────────────────────

def test_fixture_matches_bot_contract():
    assert set(op("А", "Б", local(5, 8), 60, 10)) == DASH_OPERATION_FIELDS
    assert set(op("А", "Б", local(5, 8), status="open")) == DASH_OPERATION_FIELDS
    assert set(pause(local(5, 8), 10)) == DASH_PAUSE_FIELDS
    demo = json.loads((HERE.parent / "kc_demo_data.json").read_text(encoding="utf-8"))
    assert set(demo["operations"][0]) == DASH_OPERATION_FIELDS and set(demo["pauses"][0]) == DASH_PAUSE_FIELDS
    assert {"ok", "operation_count", "pause_count", "operations", "pauses", "catalog"} <= set(demo)


def test_normalize_timezone_and_own_metrics():
    o = ks.normalize_operation(op("Эклер", "Отсадка", local(5, 8), 90, 450, 3))
    assert o["start_at"].utcoffset() == timedelta(hours=3) and o["start_at"].hour == 8
    assert o["duration_min"] == 90 and o["valid"] and o["person_hours"] == 4.5


# ── статусы ───────────────────────────────────────────────────────────────────

def test_open_closed_cancelled_counted_separately():
    ops = rum_baba() + [op("Эклер", "Отсадка", local(5, 9), status="open"),
                        op("Эклер", "Глазировка", local(5, 10), status="cancelled")]
    k = summary(snap(ops))["kpi"]
    assert (k["completed"], k["started"], k["cancelled"], k["in_work"]) == (3, 5, 1, 1)
    assert k["pcs"] == 1800  # открытая и отменённая — без штук
    assert summary(snap(ops), status="cancelled")["kpi"]["completed"] == 0
    assert summary(snap(ops), status="open")["kpi"]["in_work"] == 1


def test_open_operation_not_in_averages():
    ops = [op("Эклер", "Отсадка", local(5, 8), 60, 400, 2), op("Эклер", "Отсадка", local(5, 9), status="open")]
    s = summary(snap(ops))
    assert s["kpi"]["avg_duration"] == 60.0 and s["norms"][0]["completed"] == 1
    assert s["open_ops"][0]["operation"] == "Отсадка"


# ── человеко-часы и выработка ─────────────────────────────────────────────────

def test_person_hours_and_rates():
    s = summary(snap(rum_baba()))
    norms = {r["operation"]: r for r in s["norms"]}
    # 3 ч × 1 чел, 2 ч × 2 чел, 1 ч × 3 чел
    assert [norms[x]["person_hours"] for x in ("Накалывание", "Купание", "Декорирование")] == [3.0, 4.0, 3.0]
    assert norms["Накалывание"]["pcs_per_hour"] == 200.0 and norms["Накалывание"]["pcs_per_ph"] == 200.0
    assert norms["Купание"]["pcs_per_hour"] == 300.0 and norms["Купание"]["pcs_per_ph"] == 150.0
    assert norms["Декорирование"]["pcs_per_ph"] == 200.0
    assert s["kpi"]["person_hours"] == 10.0


def test_group_rate_is_sum_over_sum_not_mean_of_ratios():
    ops = [op("Эклер", "Отсадка", local(5, 8), 60, 100, 1),      # 100 шт/чел·ч, 1 чел·ч
           op("Эклер", "Отсадка", local(5, 10), 180, 900, 1)]    # 300 шт/чел·ч, 3 чел·ч
    n = summary(snap(ops))["norms"][0]
    assert n["pcs_per_ph"] == 250.0          # 1000 / 4, а не (100 + 300) / 2 = 200
    assert (n["min_duration"], n["max_duration"], n["avg_duration"]) == (60.0, 180.0, 120.0)


def test_invalid_duration_excluded_from_norms():
    ops = [op("Эклер", "Отсадка", local(5, 8), 60, 400, 2), op("Эклер", "Отсадка", local(5, 9), 20 * 60, 400, 2)]
    s = summary(snap(ops))
    n = s["norms"][0]
    assert n["completed"] == 2 and n["invalid"] == 1 and n["avg_duration"] == 60.0 and n["pcs_per_ph"] == 200.0
    assert s["kpi"]["invalid"] == 1 and s["kpi"]["person_hours"] == 2.0


# ── нет ложного «готового выпуска» ────────────────────────────────────────────

def test_pieces_are_not_presented_as_finished_output():
    s = summary(snap(rum_baba()))
    assert s["kpi"]["pcs"] == 1800                    # обработано по операциям — честно 1800
    assert "output" not in s["kpi"] and "output" not in json.dumps(list(s))
    labels = [t for t, _, _ in ka.kpi_rows(s["kpi"])]
    assert "Обработано изделий по операциям, шт" in labels
    assert not any("выпуск" in t.lower() and "готов" in t.lower() for t in labels)
    text = " ".join(line for _, lines in ka.build_analysis(s) for line in lines)
    assert "учтено несколько раз" in text and "не выработка готовой продукции" in text
    # в нормах изделие учтено один раз на операцию: по 600, а не 1800
    assert {r["operation"]: r["pcs"] for r in s["norms"]} == {"Накалывание": 600, "Купание": 600, "Декорирование": 600}


def test_grouping_by_product_and_operation():
    ops = rum_baba() + rum_baba(6) + [op("Торт", "Декорирование", local(6, 18), 60, 40, 2)]
    norms = {(r["product"], r["operation"]): r for r in summary(snap(ops))["norms"]}
    assert len(norms) == 4
    assert norms[("Ромовая баба", "Декорирование")]["completed"] == 2
    assert norms[("Торт", "Декорирование")]["pcs"] == 40  # та же операция, другая продукция — отдельно
    ops_by_name = {r["operation"]: r for r in summary(snap(ops))["operations"]}
    assert ops_by_name["Декорирование"]["completed"] == 3


# ── смены, полночь, период по завершению ──────────────────────────────────────

def test_midnight_operation_duration_shift_and_date():
    o = op("Эклер", "Глазировка", local(5, 23), 120, 400, 2)   # 23:00 05.10 → 01:00 06.10
    s5 = ka.build_summary(snap([o]), {"date_from": "2026-10-05", "date_to": "2026-10-05"}, now=NOW)
    s6 = ka.build_summary(snap([o]), {"date_from": "2026-10-06", "date_to": "2026-10-06"}, now=NOW)
    assert (s5["kpi"]["started"], s5["kpi"]["completed"]) == (1, 0)
    assert (s6["kpi"]["started"], s6["kpi"]["completed"], s6["kpi"]["pcs"]) == (0, 1, 400)
    assert s6["kpi"]["person_hours"] == 4.0 and s6["kpi"]["avg_duration"] == 120.0
    shift = s6["shifts"][0]
    assert (shift["shift_date"], shift["shift"], shift["midnight"]) == ("2026-10-05", "night", 1)
    assert s6["ops"][0]["midnight"] is True


def test_shift_filter_by_completion_time():
    day = op("Эклер", "Отсадка", local(5, 9), 60, 100, 1)      # закончилась в 10:00
    night = op("Эклер", "Отсадка", local(5, 19, 30), 60, 300, 1)  # закончилась в 20:30
    early = op("Эклер", "Отсадка", local(6, 6), 60, 50, 1)     # закончилась в 07:00 — ночь
    sn = snap([day, night, early])
    assert summary(sn, shift="day")["kpi"]["pcs"] == 100
    assert summary(sn, shift="night")["kpi"]["pcs"] == 350
    assert summary(sn, hour="20")["kpi"]["pcs"] == 300


def test_started_before_period_finished_inside_is_counted():
    o = op("Ромовая баба", "Накалывание", local(4, 22), 180, 600, 1)   # 04.10 22:00 → 05.10 01:00
    s = ka.build_summary(snap([o]), {"date_from": "2026-10-05", "date_to": "2026-10-11"}, now=NOW)
    assert s["kpi"]["completed"] == 1 and s["kpi"]["pcs"] == 600 and s["kpi"]["started"] == 0
    assert s["daily"][0]["date"] == "2026-10-05" and s["ops"][0]["id"] == o["id"]


def test_source_requests_extra_day_before_window():
    async def run():
        state = {"calls": [], "payload": {"ok": True, "operations": [], "pauses": [], "catalog": {}}}
        server = TestServer(fake_bot(state))
        await server.start_server()
        try:
            with patched(ks, KC_API_URL=str(server.make_url("")).rstrip("/"), KC_API_KEY=KEY, KC_SOURCE_MODE="api"):
                await ks.KcSource().get()
            asked = datetime.fromisoformat(state["calls"][0]["query"]["date_from"]).date()
            assert asked <= (datetime.now(TZ) - timedelta(days=sources.MAX_DAYS + 1)).date()
        finally:
            await server.close()
    asyncio.run(run())


# ── сотрудники ────────────────────────────────────────────────────────────────

def test_employee_filter_without_invented_attribution():
    ops = [op("Эклер", "Отсадка", local(5, 8), 60, 400, 4, user="Анна", closer="Мария"),
           op("Эклер", "Начинка", local(5, 10), 60, 300, 2, user="Ольга", closer="Ольга")]
    sn = snap(ops)
    assert summary(sn, employee="Мария")["kpi"]["pcs"] == 400      # завершила — операция её
    assert summary(sn, employee="Анна")["kpi"]["pcs"] == 400       # начала — тоже её
    assert summary(sn, employee="Анна")["kpi"]["person_hours"] == 4.0  # чел·ч группы, не делятся
    emp = {e["employee"]: e for e in summary(sn)["employees"]}
    assert (emp["Анна"]["started"], emp["Анна"]["closed"], emp["Анна"]["pcs_closed"]) == (1, 0, 0)
    assert (emp["Мария"]["started"], emp["Мария"]["closed"], emp["Мария"]["pcs_closed"]) == (0, 1, 400)
    assert not any(k in emp["Мария"] for k in ("pcs_per_ph", "person_hours", "pcs_per_person"))
    assert emp["Мария"]["hours"][9] == 1 and emp["Анна"]["hours"][8] == 1


# ── критические остановки ─────────────────────────────────────────────────────

def test_pauses_count_duration_reasons_and_filters():
    ps = [pause(local(5, 10), 20, "Нет сырья", "Анна"), pause(local(5, 15), 40, "Поломка оборудования", "Ольга"),
          pause(local(6, 23, 50), 30, "Нет сырья", "Анна")]   # через полночь: 30 мин
    s = summary(snap(rum_baba(), ps))
    assert (s["kpi"]["pauses"], s["kpi"]["stop_minutes"]) == (3, 90.0)
    assert {r["label"]: (r["count"], r["value"]) for r in s["pause_reasons"]} == {
        "Нет сырья": (2, 50.0), "Поломка оборудования": (1, 40.0)}
    assert [d["minutes"] for d in s["pause_daily"]] == [60.0, 30.0]
    assert summary(snap(rum_baba(), ps), employee="Ольга")["kpi"]["pauses"] == 1
    assert summary(snap(rum_baba(), ps), product="Ромовая баба")["kpi"]["pauses"] == 3  # нет продукции у остановки


def test_lfl_and_norm_change_vs_previous_period():
    prev = [op("Эклер", "Отсадка", local(3, 8), 60, 200, 1)]   # 200 шт/чел·ч
    cur = [op("Эклер", "Отсадка", local(12, 8), 60, 300, 1)]   # 300 шт/чел·ч
    s = ka.build_summary(snap(prev + cur), {"date_from": "2026-10-09", "date_to": "2026-10-15"}, now=NOW)
    assert s["lfl"]["available"] and s["lfl"]["kpi"]["pcs"] == {"prev": 200, "change": 50.0}
    n = s["norms"][0]
    assert (n["prev_pcs_per_ph"], n["pcs_per_ph_change"]) == (200.0, 50.0)


def test_no_data_is_shown_as_no_data():
    k = summary(snap([]))["kpi"]
    rows = {t: v for t, v, _ in ka.kpi_rows(k)}
    assert rows["Обработано изделий по операциям, шт"] == "нет данных"
    assert rows["Средняя выработка, шт/чел·ч"] == "нет данных" and k["pcs_per_ph"] is None


# ── API: ключ, недоступность ──────────────────────────────────────────────────

def fake_bot(state):
    async def all_(request):
        state["calls"].append({"query": dict(request.query), "key": request.headers.get("X-API-Key")})
        if request.headers.get("X-API-Key") != KEY:
            return web.json_response({"error": "unauthorized"}, status=401)
        if state.get("fail"):
            return web.json_response({"error": "server_error"}, status=500)
        return web.json_response(state["payload"])

    app = web.Application()
    app.router.add_get("/api/dashboard/all", all_)
    return app


class patched:
    def __init__(self, module, **values):
        self.module, self.values, self.old = module, values, {}

    def __enter__(self):
        for k, v in self.values.items():
            self.old[k] = getattr(self.module, k)
            setattr(self.module, k, v)

    def __exit__(self, *exc):
        for k, v in self.old.items():
            setattr(self.module, k, v)


def test_api_key_header_errors_and_cache():
    async def run():
        state = {"calls": [], "payload": {"ok": True, "operation_count": 3, "pause_count": 0, "operations": rum_baba(),
                                          "pauses": [], "catalog": {"reasons": []}}}
        server = TestServer(fake_bot(state))
        await server.start_server()
        url = str(server.make_url("")).rstrip("/")
        try:
            with patched(ks, KC_API_URL=url, KC_API_KEY=KEY, KC_SOURCE_MODE="api"):
                src = ks.KcSource()
                first = await src.get()
                assert len(first["operations"]) == 3
                call = state["calls"][0]
                assert call["key"] == KEY and "api_key" not in call["query"]
                assert KEY not in json.dumps(src.describe())
                state["fail"] = True
                again = await src.get(force=True)         # бот упал — остаётся последний снимок
                assert again is first and "500" in src.describe()["error"]
            with patched(ks, KC_API_URL=url, KC_API_KEY="wrong", KC_SOURCE_MODE="api"):
                try:
                    await ks.KcSource().get()
                    raise AssertionError("ожидалась ошибка ключа")
                except RuntimeError as exc:
                    assert "401" in str(exc)
            with patched(ks, KC_API_URL=url, KC_API_KEY="", KC_SOURCE_MODE="api"):
                try:
                    await ks.KcSource().get()
                    raise AssertionError("без ключа запрос не уходит")
                except RuntimeError as exc:
                    assert "KC_API_KEY" in str(exc)
        finally:
            await server.close()
    asyncio.run(run())


async def _client():
    app = dashboard.build_app()
    app.on_startup.clear()
    client = TestClient(TestServer(app))
    await client.start_server()
    return client


def _fresh_kc():
    dashboard.kc_source = ks.KcSource()
    dashboard.SOURCES["kc"]["data"] = dashboard.kc_source


def test_kc_down_does_not_break_other_sections():
    async def run():
        _fresh_kc()
        client = await _client()
        try:
            with patched(ks, KC_API_URL="http://127.0.0.1:9", KC_API_KEY=KEY, KC_SOURCE_MODE="api", KC_TIMEOUT=3):
                for path in ("/api/summary?source=kc", "/api/meta?source=kc", "/api/export/pdf?source=kc",
                             "/api/export/xlsx?source=kc"):
                    r = await client.get(path)
                    body = await r.text()
                    assert r.status == 503 and "API кондитерского цеха недоступно" in json.loads(body)["error"], path
                    assert KEY not in body
                for src in ("fk", "bhm", "proofing"):
                    r = await client.get(f"/api/summary?source={src}")
                    assert r.status == 200 and (await r.json())["ok"], src
                h = await client.get("/health")
                body = await h.json()
                assert h.status == 200 and "kc" in body and KEY not in json.dumps(body)
                meta = await (await client.get("/api/meta?source=fk")).json()
                assert {k: meta["sources"][k] for k in ("fk", "bhm", "proofing", "kc")} == {
                    "fk": True, "bhm": True, "proofing": True, "kc": True}
            with patched(ks, KC_SOURCE_MODE="off"):
                r = await client.get("/api/summary?source=kc")
                assert r.status == 503 and "не настроен" in (await r.json())["error"]
        finally:
            await client.close()
    asyncio.run(run())


def test_basic_auth_protects_kc():
    async def run():
        client = await _client()
        try:
            with patched(dashboard, RELAXDEV_DEPLOY=True, DASHBOARD_USER="u", DASHBOARD_PASSWORD="p"):
                for path in ("/api/summary?source=kc", "/api/export/xlsx?source=kc", "/api/meta?source=kc"):
                    assert (await client.get(path)).status == 401, path
        finally:
            await client.close()
    asyncio.run(run())


# ── Excel и PDF ───────────────────────────────────────────────────────────────

def demo_summary():
    d = json.loads((HERE.parent / "kc_demo_data.json").read_text(encoding="utf-8"))
    s = ka.build_summary(ks.normalize(d["operations"], d["pauses"], d["catalog"]),
                         {"date_from": "2026-10-05", "date_to": "2026-10-11"}, now=NOW)
    return s, [("Источник", "КЦ — кондитерский цех"), ("Период", "2026-10-05 – 2026-10-11")]


def test_excel_all_widgets_numbers_and_notes():
    import openpyxl
    s, filters = demo_summary()
    wb = openpyxl.load_workbook(io.BytesIO(ke.build_xlsx(s, "all", filters)))
    assert wb.sheetnames == [ke.widget_rows(s, w)[0][:31] for w in ke.ALL_WIDGETS] + ["Фильтры"]
    daily = list(wb["По дням"].iter_rows(min_row=2, values_only=True))
    assert sum(r[4] for r in daily) == s["kpi"]["pcs"] and sum(r[3] for r in daily) == s["kpi"]["completed"]
    norms = list(wb["Нормы"].iter_rows(min_row=2, values_only=True))
    assert len(norms) == len(s["norms"]) and sum(r[3] for r in norms) == s["kpi"]["pcs"]
    notes = {r[0]: r[1] for r in wb["Фильтры"].iter_rows(min_row=2, values_only=True) if r[0]}
    assert "учтено несколько раз" in notes["Обработано изделий"] and "не утверждённый" in notes["Нормы"]
    kpi = {r[0]: r[1] for r in wb["KPI"].iter_rows(min_row=2, values_only=True)}
    assert kpi["Обработано изделий по операциям, шт"] == ka.fmt_int(s["kpi"]["pcs"])
    assert wb["Операции"].max_row - 1 == len(s["ops"])


def test_pdf_builds():
    s, filters = demo_summary()
    blob = ke.build_pdf(s, ka.build_analysis(s, filters), filters)
    assert blob[:4] == b"%PDF" and len(blob) > 10_000


def test_app_exports_use_same_filters_as_page():
    async def run():
        import openpyxl
        with patched(ks, KC_SOURCE_MODE="demo"):
            _fresh_kc()
            client = await _client()
            try:
                params = {"source": "kc", "date_from": "2026-10-05", "date_to": "2026-10-11", "operation": "Сборка",
                          "shift": "day"}
                page = await (await client.get("/api/summary", params=params)).json()
                assert page["ok"] and {r["operation"] for r in page["norms"]} <= {"Сборка"}
                r = await client.get("/api/export/xlsx", params={**params, "widget": "kpi"})
                assert r.status == 200 and "kc_kpi_2026-10-05_2026-10-11" in r.headers["Content-Disposition"]
                wb = openpyxl.load_workbook(io.BytesIO(await r.read()))
                kpi = {row[0]: row[1] for row in wb["KPI"].iter_rows(min_row=2, values_only=True)}
                assert kpi["Завершено операций"] == ka.fmt_int(page["kpi"]["completed"])
                filt = {row[0]: row[1] for row in wb["Фильтры"].iter_rows(min_row=2, values_only=True) if row[0]}
                assert filt["Операция"] == "Сборка" and filt["Смена"].startswith("дневная")
                r = await client.get("/api/export/pdf", params=params)
                assert r.status == 200 and (await r.read())[:4] == b"%PDF"
                r = await client.get("/api/export/xlsx", params={**params, "widget": "plan"})
                assert r.status == 400   # плана у КЦ нет
            finally:
                await client.close()
    asyncio.run(run())


# ── остальные разделы не изменились ───────────────────────────────────────────

FK_SUMMARY_KEYS = {"kpi", "daily", "hourly", "hourly_target", "hourly_below", "plan", "by_equipment",
                   "top_products", "pause_reasons", "work_vs_stop", "top_employees", "product_timings",
                   "equipment_timings", "employee_timings", "detail", "lfl", "source", "source_key",
                   "source_title", "units", "plan_enabled", "excluded_dates", "ok", "filters", "analysis"}


def test_fk_bhm_proofing_unchanged():
    async def run():
        client = await _client()
        try:
            for src, unit in (("fk", "шт"), ("bhm", "кг")):
                data = await (await client.get(f"/api/summary?source={src}")).json()
                assert set(data) == FK_SUMMARY_KEYS and data["units"]["qty"] == unit
            assert (await (await client.get("/api/summary")).json())["source_key"] == "fk"
            pr = await (await client.get("/api/summary?source=proofing&date_from=2026-10-05&date_to=2026-10-11")).json()
            assert pr["ok"] and pr["kind"] == "proofing" and pr["kpi"]["output"] > 0
            for src in ("fk", "bhm", "proofing"):
                r = await client.get(f"/api/export/pdf?source={src}")
                assert r.status == 200 and (await r.read())[:4] == b"%PDF", src
        finally:
            await client.close()
    asyncio.run(run())


if __name__ == "__main__":
    tests = [v for k, v in globals().items() if k.startswith("test_")]
    for t in tests:
        t()
        print(f"  ✅ {t.__name__}")
    print(f"ИТОГ: {len(tests)} тестов прошли")
