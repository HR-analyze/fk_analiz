"""Справочник соответствий для Ганта: продукты цехов → единый продукт, операции → единая операция.

У ботов свои названия («Батон / ромашка» в БХМ, «Батон классический» на формовке), а в КЦ
операции вводят руками («Начинка», «Внесение начинки», «Набивка начинкой»). Чтобы собрать
путь продукта БХМ → формовка → расстойка → отпекание → КЦ и среднее по продукту, названия
сводятся к одному. Справочник загружают Excel-файлом на дашборде.

Без справочника названия сводятся только автоматически: регистр, «ё», пробелы и знаки
препинания не различаются. Похожие названия (опечатки, «фреш») только подсказываются —
склеивать их или нет, решает человек в справочнике.

Хранение — там же, где планы ФК, но в ОТДЕЛЬНОЙ таблице fk_dashboard_product_map:
таблицу планов (fk_dashboard_plan_state) этот модуль не читает и не меняет.
Без PostgreSQL (локальный запуск) — JSON-файл рядом с PLAN_FILE.
"""
import io
import json
import os
import re
import tempfile
from difflib import SequenceMatcher
from datetime import datetime
from pathlib import Path
from threading import Lock

import plans

SECTIONS = {"bhm": "БХМ", "fk": "Формовка", "proofing": "Расстойка", "kc": "КЦ"}
# «ФК» — прежнее название цеха в справочнике: старые файлы продолжают загружаться
TITLE_TO_KEY = {v.lower(): k for k, v in SECTIONS.items()} | {k: k for k in SECTIONS} | {"фк": "fk"}
OPS_SHEET = "Операции"
SIMILAR_RATIO = 0.8
SIMILAR_OPS_RATIO = 0.86  # операции короткие: «Подготовка форм» и «Подготовка крема» — не опечатка
MAX_ROWS = 5000
_lock = Lock()


class MapError(Exception):
    """Понятная пользователю ошибка разбора файла."""


def _file():
    return Path(os.getenv("PRODUCT_MAP_FILE", str(Path(plans.PLAN_FILE).with_name("product_map.json"))))


def _pg():
    import pg_plan_store
    return pg_plan_store._connect()


def init_storage():
    if not plans.PG_PLANS:
        return
    with _pg() as db:
        db.execute("""
            CREATE TABLE IF NOT EXISTS fk_dashboard_product_map (
                id smallint PRIMARY KEY CHECK (id=1),
                payload jsonb NOT NULL DEFAULT '{}'::jsonb,
                updated_at timestamptz NOT NULL DEFAULT now()
            )
        """)
        db.execute("INSERT INTO fk_dashboard_product_map(id,payload) VALUES (1,'{}'::jsonb) ON CONFLICT (id) DO NOTHING")


def _clean(name):
    return re.sub(r"\s+", " ", str(name or "")).strip()


def norm(name):
    """Ключ сравнения: без регистра, «ё» = «е», знаки препинания и лишние пробелы не важны."""
    s = _clean(name).lower().replace("ё", "е")
    return re.sub(r"\s+", " ", re.sub(r"[^\w]+", " ", s)).strip()


def similar(names, ratio=SIMILAR_RATIO):
    """{name: [похожие]} — подсказка для справочника; одинаковые после norm не считаются похожими."""
    keys = sorted({norm(n): n for n in names}.items())
    out = {}
    sm = SequenceMatcher(autojunk=False)
    for i, (ka, a) in enumerate(keys):
        sm.set_seq2(ka)  # вторая строка кешируется — сравниваем одну со всеми
        for kb, b in keys[i + 1:]:
            if 2 * min(len(ka), len(kb)) / (len(ka) + len(kb)) < ratio:  # верхняя граница ratio()
                continue
            sm.set_seq1(kb)
            if sm.real_quick_ratio() >= ratio and sm.quick_ratio() >= ratio and sm.ratio() >= ratio:
                out.setdefault(a, []).append(b)
                out.setdefault(b, []).append(a)
    return out


def load():
    """{"map": {section: {product: unified}}, "ops": {section: {operation: unified}}, "updated_at", "file"}."""
    if plans.PG_PLANS:
        with _pg() as db:
            row = db.execute("SELECT payload FROM fk_dashboard_product_map WHERE id=1").fetchone()
        payload = (row[0] if row else None) or {}
    else:
        path = _file()
        payload = json.loads(path.read_text(encoding="utf-8")) if path.exists() else {}
    payload.setdefault("map", {})
    payload.setdefault("ops", {})
    return payload


def save(mapping, filename="", ops=None):
    payload = {"map": mapping, "ops": ops or {}, "updated_at": datetime.now().astimezone().isoformat(), "file": filename}
    with _lock:
        if plans.PG_PLANS:
            from psycopg.types.json import Jsonb
            with _pg() as db:
                db.execute("UPDATE fk_dashboard_product_map SET payload=%s, updated_at=now() WHERE id=1", (Jsonb(payload),))
        else:
            path = _file()
            path.parent.mkdir(parents=True, exist_ok=True)
            fd, tmp = tempfile.mkstemp(dir=str(path.parent), suffix=".tmp")
            with os.fdopen(fd, "w", encoding="utf-8") as fh:
                json.dump(payload, fh, ensure_ascii=False, indent=1)
            os.replace(tmp, path)
    return payload


def unify(mapping, section, name):
    """Единое имя; нет в справочнике — само название. Строка справочника находится и при
    другом регистре, «ё» или пробелах (руками вводят по-разному)."""
    items = mapping.get(section) or {}
    if name in items:
        return items[name]
    key = norm(name)
    for src, unified in items.items():
        if norm(src) == key:
            return unified
    return name


def template_xlsx(products, mapping, operations=None, ops_map=None):
    """Шаблон: лист «Соответствия» (продукты всех цехов) и «Операции». Колонка «Единый …» —
    текущая; «Похоже на» — подсказка (опечатки, «фреш»), на загрузку не влияет."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill

    wb = openpyxl.Workbook()
    head_fill = PatternFill("solid", fgColor="E8EEFB")

    def head(ws, titles, widths):
        ws.append(titles)
        for cell in ws[1]:
            cell.font, cell.fill = Font(bold=True), head_fill
        for col, w in zip("ABCDE", widths):
            ws.column_dimensions[col].width = w
        ws.freeze_panes = "A2"

    ws = wb.active
    ws.title = "Соответствия"
    head(ws, ["Цех", "Продукт в цехе", "Единый продукт", "Похоже на (подсказка)"], (12, 46, 46, 60))
    hints = similar([n for names in products.values() for n in names])
    for key, title in SECTIONS.items():
        for name in sorted(products.get(key, []), key=str.lower):
            ws.append([title, name, (mapping.get(key) or {}).get(name, ""), "; ".join(sorted(hints.get(name, [])))])
    ops_ws = wb.create_sheet(OPS_SHEET)
    head(ops_ws, ["Цех", "Операция в цехе", "Единая операция", "Похоже на (подсказка)", "Где встречается"], (12, 40, 40, 50, 70))
    for key, title in SECTIONS.items():
        ops = (operations or {}).get(key) or {}
        op_hints = similar(list(ops), SIMILAR_OPS_RATIO)
        for name in sorted(ops, key=str.lower):
            ops_ws.append([title, name, ((ops_map or {}).get(key) or {}).get(name, ""),
                           "; ".join(sorted(op_hints.get(name, []))), ", ".join(sorted(ops[name], key=str.lower))[:300]])
    help_ws = wb.create_sheet("Как заполнять")
    for line in ("Лист «Соответствия»: в колонке «Единый продукт» напиши одно название продукта для всех цехов.",
                 "Пример: БХМ «Батон / ромашка» и Формовка «Батон классический» → «Батон классический».",
                 "Лист «Операции»: одно название для одной операции, например «Начинка» и «Набивка начинкой» → «Внесение начинки».",
                 "Пустая ячейка — название остаётся как есть. Регистр, «ё», пробелы и знаки препинания сводятся автоматически.",
                 "«Похоже на» — только подсказка (опечатки, «фреш»): склеивать или нет, решаешь ты.",
                 "Загружай файл на вкладке «Справочник продуктов» раздела «Гант» с ключом загрузки плана."):
        help_ws.append([line])
    help_ws.column_dimensions["A"].width = 110
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _parse_sheet(ws, name_cols, unified_col, what):
    rows = ws.iter_rows(values_only=True)
    head = [(_clean(c) or "").lower() for c in next(rows, [])]

    def col(*names):
        for i, h in enumerate(head):
            if any(n in h for n in names):
                return i
        return None

    i_sec, i_name, i_unified = col("цех"), col(*name_cols), col(unified_col)
    if None in (i_sec, i_name, i_unified):
        raise MapError(f"{what}: нужны колонки «Цех», «{name_cols[0].capitalize()}», «{unified_col.capitalize()} …» — скачай шаблон")
    mapping, count, unknown = {}, 0, set()
    for n, row in enumerate(rows, 2):
        if n > MAX_ROWS:
            raise MapError(f"{what}: больше {MAX_ROWS} строк")
        sec = _clean(row[i_sec] if i_sec < len(row) else "").lower()
        name = _clean(row[i_name] if i_name < len(row) else "")
        unified = _clean(row[i_unified] if i_unified < len(row) else "")
        if not sec and not name:
            continue
        key = TITLE_TO_KEY.get(sec)
        if not key:
            unknown.add(sec)
            continue
        if name and unified and unified != name:
            mapping.setdefault(key, {})[name] = unified
            count += 1
    if unknown:
        raise MapError(f"{what}: неизвестный цех: " + ", ".join(sorted(unknown)) + " (нужно: " + ", ".join(SECTIONS.values()) + ")")
    return mapping, count


def parse_workbook(stream):
    """{"map", "ops", "pairs", "op_pairs"}; лист «Операции» необязателен (старые файлы — только продукты)."""
    import openpyxl
    try:
        wb = openpyxl.load_workbook(stream, read_only=True, data_only=True)
    except Exception as exc:
        raise MapError(f"не удалось открыть файл как Excel: {exc}") from exc
    product_ws = next((ws for ws in wb.worksheets if ws.title != OPS_SHEET), None)
    if product_ws is None:
        raise MapError("нет листа с продуктами — скачай шаблон")
    mapping, pairs = _parse_sheet(product_ws, ("продукт в цехе", "в цехе"), "единый", "продукты")
    ops, op_pairs = {}, 0
    if OPS_SHEET in wb.sheetnames:
        ops, op_pairs = _parse_sheet(wb[OPS_SHEET], ("операция в цехе", "в цехе"), "единая", "операции")
    return {"map": mapping, "ops": ops, "pairs": pairs, "op_pairs": op_pairs}


def parse_xlsx(stream):
    """Только продукты: (mapping, count)."""
    parsed = parse_workbook(stream)
    return parsed["map"], parsed["pairs"]
