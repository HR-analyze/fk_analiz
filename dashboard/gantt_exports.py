"""Выгрузки раздела «Гант»: Excel (операции, общий тайминг, передачи на фасовку, справочник)
и PDF с диаграммой. Берут тот же расчёт gantt.build, что и страница."""
import io
from datetime import datetime

from exports import TZ, _register_fonts

SECTION_COLORS = {"bhm": "#EB6834", "fk": "#2A78D6", "proofing": "#1BAF7A", "kc": "#7C5CF5"}
STAGE_COLORS = {"bhm": "#EB6834", "fk": "#2A78D6", "proofing": "#1BAF7A", "baking": "#0F766E", "kc": "#7C5CF5",
                "pack": "#9CA3AF"}
FLAG_NAMES = {"ok": "", "check": "проверить: время по сообщению", "invalid": "некорректное время"}


def _v(value):
    return "" if value is None else value


def _hm(minutes):
    if minutes is None:
        return ""
    m = int(round(minutes))
    sign = "−" if m < 0 else ""
    m = abs(m)
    return f"{sign}{m} мин" if m < 60 else f"{sign}{m // 60} ч {m % 60:02d} мин"


def bar_rows(summary):
    rows = []
    for r in summary["rows"]:
        for b in r["bars"]:
            rows.append([b["section_title"], r["product"], b["product"] if b["product"] != r["product"] else "",
                         b["operation"], b["start"][:16].replace("T", " "),
                         b["end"][:16].replace("T", " ") if b["end"] else ("идёт" if b["running"] else "конец не отмечен"),
                         _v(b["duration_min"]), _v(b["people"]), _v(b["person_hours"]), _v(b["qty"]), b["unit"],
                         _v(b["rate"]), _v(b["rate_per_person"]), b["equipment"], b["user"], b["closed_by"],
                         "чат Telegram" if b["source"] == "chat" else "бот", FLAG_NAMES.get(b["flag"], b["flag"]), b["note"]])
    return rows


BAR_COLUMNS = ["Цех", "Продукт (единый)", "Название в цехе", "Операция", "Начало", "Конец", "Длительность, мин",
               "Людей", "Чел·ч", "Выработка", "Ед.", "В час", "На человека в час", "Оборудование", "Начал",
               "Завершил", "Источник", "Отметка", "Примечание"]


def timing_rows(summary):
    rows = []
    for t in summary["timing"]:
        row = [t["product"], t["stages_count"], _hm(t["lead_min"])]
        for st in t["stages"]:
            if st["key"] == "pack":
                hs = st.get("handoffs") or []
                row.append("; ".join(f"передано {h['qty']} {h['unit']} в {h['time']}" for h in hs) or "нет источника")
            elif st.get("present"):
                row.append(f"{st['start_time']}–{st['end_time'] or '…'} · {_hm(st['span_min']) or '—'}"
                           f" · {st['qty']}" + (f" · до {st['people_max']} чел." if st.get("people_max") else "")
                           + (f" · ожидание {_hm(st['gap_min'])}" if st.get("gap_min") is not None else ""))
            else:
                row.append("—")
        rows.append(row)
    return rows


def timing_columns(summary):
    return ["Продукт (единый)", "Этапов с данными", "От первого начала до последнего конца"] + [s["title"] for s in summary["stages"]]


def build_analysis(summary, filters=None):
    t = summary["totals"]
    rows = [f"Окно: {summary['day']}, {summary['window_title']}.",
            f"Операций на шкале: {t['ops']} (закрыто {t['closed']}, без отметки конца или идут {t['open']}).",
            "По цехам: " + ", ".join(f"{k} — {v}" for k, v in t["by_section"].items()) + ".",
            f"С известным числом людей: {t['people_known']}; человеко-часов (где известно): {str(t['person_hours']).replace('.', ',')}."]
    if t["from_chat"]:
        rows.append(f"Из разобранного чата КЦ: {t['from_chat']} (до запуска бота КЦ); "
                    f"время по времени сообщения — {t['check']}.")
    blocks = [("Итоги окна", rows)]
    chains = [x for x in summary["timing"] if x["stages_count"] > 1]
    if chains:
        blocks.append(("Общий тайминг: продукты, прошедшие несколько цехов", [
            f"{x['product']}: этапов {x['stages_count']}, от первого начала до последнего конца {_hm(x['lead_min']) or '—'}."
            for x in chains[:10]]))
    else:
        blocks.append(("Общий тайминг", ["В этом окне нет продукта, найденного больше чем в одном цехе. "
                                         "Сведи названия в справочнике продуктов, если у цехов они разные."]))
    if summary["handoffs"]:
        blocks.append(("Передано на фасовку (из чата КЦ)", [f"{h['product']} — {h['qty']} {h['unit']} в {h['time']}"
                                                           for h in summary["handoffs"]]))
    errors = summary.get("section_errors") or {}
    if errors:
        blocks.append(("Источники с ошибками", [f"{k}: {v}" for k, v in errors.items()]))
    blocks.append(("Оговорки", ["Цеха связаны по единому продукту, а не по номеру партии — у ботов нет общего номера.",
                                "Фасовки пока нет ни в одном боте: этап заложен, данных нет.",
                                "Людей знают только КЦ (бот и чат); у ФК, БХМ и расстойки — отметивший оператор."]))
    if filters:
        blocks.append(("Применённые фильтры", [f"{n}: {v}" for n, v in filters if v]))
    return blocks


def build_xlsx(summary, filters, mapping):
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    head_fill, head_font = PatternFill("solid", fgColor="E8EEFB"), Font(bold=True)

    def sheet(title, columns, rows, widths=None):
        ws = wb.create_sheet(title[:31])
        ws.append(columns)
        for row in rows:
            ws.append(row)
        if not rows:
            ws.append(["нет данных"])
        for i, c in enumerate(columns, 1):
            cell = ws.cell(row=1, column=i)
            cell.fill, cell.font = head_fill, head_font
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            longest = max([len(str(c))] + [len(str(r[i - 1])) for r in rows[:200] if i - 1 < len(r)])
            ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(i, min(max(longest + 2, 9), 60))
        ws.freeze_panes = "A2"
        if rows:
            ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{len(rows) + 1}"

    sheet("Операции", BAR_COLUMNS, bar_rows(summary))
    sheet("Общий тайминг", timing_columns(summary), timing_rows(summary))
    sheet("Передано на фасовку", ["Продукт", "Когда", "Количество", "Ед.", "Время точное"],
          [[h["product"], h["time"], h["qty"], h["unit"], "да" if h["exact"] else "по сообщению"] for h in summary["handoffs"]])
    sheet("Справочник продуктов", ["Цех", "Продукт в цехе", "Единый продукт"],
          [[{"bhm": "БХМ", "fk": "ФК", "proofing": "Расстойка", "kc": "КЦ"}.get(sec, sec), name, unified]
           for sec, items in sorted(mapping.items()) for name, unified in sorted(items.items())])
    ws = wb.create_sheet("Фильтры")
    ws.append(["Параметр", "Значение"])
    for n, v in filters:
        ws.append([n, v])
    ws.append(["Выгружено", datetime.now(TZ).strftime("%d.%m.%Y %H:%M")])
    for line in build_analysis(summary)[-1][1]:
        ws.append(["Оговорка", line])
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 26, 100
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _gantt_drawing(summary, rows, width, font, row_h=11):
    """Диаграмма на примитивах reportlab: строка — продукт · операция, полосы по времени."""
    from reportlab.graphics.shapes import Drawing, Line, Rect, String
    from reportlab.lib.colors import HexColor

    start = datetime.fromisoformat(summary["start"])
    end = datetime.fromisoformat(summary["end"])
    total = (end - start).total_seconds() or 1
    label_w = 190
    plot_w = width - label_w
    lanes = sum(r["lanes"] for r in rows)
    height = 18 + lanes * row_h + 4
    d = Drawing(width, height)
    hours = int(total // 3600)
    for h in range(hours + 1):
        x = label_w + plot_w * h * 3600 / total
        d.add(Line(x, 0, x, height - 14, strokeColor=HexColor("#E6EAF0"), strokeWidth=0.4))
        label = (start.hour + h) % 24
        d.add(String(x, height - 10, f"{label:02d}", fontName=font, fontSize=5.5, textAnchor="middle",
                     fillColor=HexColor("#8A94A6")))
    y = height - 18
    for r in rows:
        block_h = r["lanes"] * row_h
        y -= block_h
        text = f"{r['section_title']} · {r['product']} · {r['operation']}"
        d.add(String(2, y + block_h - 8, text[:52], fontName=font, fontSize=5.6, fillColor=HexColor("#3C4658")))
        d.add(Line(0, y, width, y, strokeColor=HexColor("#F1F4F8"), strokeWidth=0.3))
        for b in r["bars"]:
            s = max(datetime.fromisoformat(b["start"]), start)
            by = y + (r["lanes"] - 1 - b["lane"]) * row_h + 2
            color = HexColor(SECTION_COLORS.get(b["section"], "#667085"))
            x = label_w + plot_w * (s - start).total_seconds() / total
            if b["draw_end"]:
                e = min(datetime.fromisoformat(b["draw_end"]), end)
                w = max(plot_w * (e - s).total_seconds() / total, 1.2)
                d.add(Rect(x, by, w, row_h - 4, fillColor=color, strokeColor=None,
                           fillOpacity=0.45 if b["flag"] != "ok" or b["running"] else 1))
                info = " · ".join(x for x in (f"{b['people']} чел" if b["people"] else "",
                                               f"{b['qty']:,.0f} {b['unit']}".replace(",", " ") if b["qty"] else "") if x)
                if info and w > 40:
                    d.add(String(x + 2, by + 2, info[:int(w / 3)], fontName=font, fontSize=4.8, fillColor=HexColor("#FFFFFF")))
            else:
                d.add(Rect(x - 1.5, by, 3, row_h - 4, fillColor=color, strokeColor=None))
    return d


def build_pdf(summary, analysis, filters, title="Диаграмма Ганта и общий тайминг"):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import KeepTogether, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    font, bold = _register_fonts()
    out = io.BytesIO()
    doc = SimpleDocTemplate(out, pagesize=landscape(A4), leftMargin=10 * mm, rightMargin=10 * mm,
                            topMargin=10 * mm, bottomMargin=10 * mm, title=title)
    width = doc.width
    h1 = ParagraphStyle("h1", fontName=bold, fontSize=16, leading=20, spaceAfter=4)
    h2 = ParagraphStyle("h2", fontName=bold, fontSize=11, leading=14, spaceBefore=6, spaceAfter=4)
    body = ParagraphStyle("body", fontName=font, fontSize=8.2, leading=11)
    muted = ParagraphStyle("muted", fontName=font, fontSize=7.6, leading=10, textColor=colors.HexColor("#667085"))
    cell = ParagraphStyle("c", fontName=font, fontSize=6.2, leading=7.6)
    story = [Paragraph(title, h1), Paragraph(", ".join(f"{n}: {v}" for n, v in filters if v), muted),
             Paragraph(f"Сформировано {datetime.now(TZ).strftime('%d.%m.%Y %H:%M')}", muted), Spacer(1, 3 * mm)]
    for heading, lines in analysis:
        story.append(KeepTogether([Paragraph(heading, h2)] + [Paragraph(x, body) for x in lines]))
    story.append(PageBreak())
    story.append(Paragraph(f"Диаграмма Ганта · {summary['day']} · {summary['window_title']}", h2))
    legend = " · ".join(f"<font color='{c}'>■</font> {n}" for n, c in (("БХМ", SECTION_COLORS["bhm"]), ("ФК", SECTION_COLORS["fk"]),
                                                                        ("Расстойка/отпекание", SECTION_COLORS["proofing"]),
                                                                        ("КЦ", SECTION_COLORS["kc"])))
    story.append(Paragraph(legend + " · бледная полоса — время по сообщению или операция идёт; штрих — конец не отмечен", muted))
    rows, chunk, lanes = summary["rows"], [], 0
    for r in rows:
        if lanes + r["lanes"] > 40 and chunk:
            story += [_gantt_drawing(summary, chunk, width, font), PageBreak()]
            chunk, lanes = [], 0
        chunk.append(r)
        lanes += r["lanes"]
    if chunk:
        story.append(_gantt_drawing(summary, chunk, width, font))
    if not rows:
        story.append(Paragraph("В этом окне операций нет.", body))
    story.append(PageBreak())

    def table(columns, data, limit=300):
        t = Table([[Paragraph(str(c), ParagraphStyle("h", fontName=bold, fontSize=6.4, leading=7.8)) for c in columns]]
                  + [[Paragraph(str(c), cell) for c in row] for row in data[:limit]], repeatRows=1)
        t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#E6EAF0")),
                               ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#EFF3FB")),
                               ("VALIGN", (0, 0), (-1, -1), "TOP"),
                               ("TOPPADDING", (0, 0), (-1, -1), 2), ("BOTTOMPADDING", (0, 0), (-1, -1), 2)]))
        return t

    story.append(Paragraph("Общий тайминг по продуктам", h2))
    story.append(table(timing_columns(summary), timing_rows(summary)))
    story.append(PageBreak())
    story.append(Paragraph("Операции", h2))
    cols = [0, 1, 3, 4, 5, 6, 7, 9, 10, 12, 13, 17]
    story.append(table([BAR_COLUMNS[i] for i in cols], [[r[i] for i in cols] for r in bar_rows(summary)]))
    doc.build(story)
    return out.getvalue()
